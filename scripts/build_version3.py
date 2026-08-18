#!/usr/bin/env python3
"""Build Version 3 of the atlas workbook from the reviewed Version 2 base.

Version 3 adds a transparent orientation layer for every corpus source. These
rows are graph coordinates inferred from the existing editorial contribution
notes; they are deliberately not promoted to canonical source entries.
"""

from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import tempfile
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.table import Table, TableStyleInfo

from reproducible import normalize_xlsx, source_fingerprint, target_for_treatment


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = ROOT / "fantasy_high_fantasy_archetype_atlas_v2.xlsx"
DEFAULT_OUTPUT = ROOT / "fantasy_high_fantasy_archetype_atlas_v3.xlsx"

NAVY = "0B1222"
NAVY_2 = "111C30"
HEADER = "17243B"
INK = "EAF1FB"
MUTED = "A7B5C9"
GOLD = "F2B35D"
TEAL = "66D9C6"
VIOLET = "A78BFA"
RED = "FB7185"
GRID = "2B3A55"


# Rules intentionally favor broad, defensible coordinates. A matching rule says
# that a corpus orientation phrase is usefully compared along this axis; it does
# not claim exact equivalence.
CONCEPT_RULES: list[tuple[str, str, str]] = [
    (r"\bdemigods?\b|divine[- ]blood", "PPL-402", "high"),
    (r"\bgods?\b|\bdeities\b|divine beings?|\bprimals?\b", "PPL-400", "moderate"),
    (r"\bangels?\b|celestials?", "PPL-404", "high"),
    (r"\bdevils?\b", "PPL-502", "high"),
    (r"demon lords?|archfiends?", "PPL-509", "high"),
    (r"\bdemons?\b|fiends?", "PPL-501", "high"),
    (r"\bjinn\b|\bdjinn\b", "PPL-203", "high"),
    (r"\bspirits?\b|spirit ecology|spirit world", "PPL-300", "moderate"),
    (r"household (beings|spirits)", "PPL-205", "high"),
    (r"forest beings?|nature spirits?|animal gods?", "PPL-200", "moderate"),
    (r"\bnymphs?\b|place-bound", "PPL-204", "high"),
    (r"\bfae\b|\bfaeries\b|\bfairies\b|fairy-tale beings?", "PPL-200", "moderate"),
    (r"\by[ōo]kai\b|\bkami\b", "PPL-200", "moderate"),
    (r"\boni\b", "PPL-501", "moderate"),
    (r"\btengu\b|\bkitsune\b|\bgumiho\b|shapeshifters?", "PPL-1000", "moderate"),
    (r"\btsukumogami\b|artificial (bodies|life)|construct(ed)? people", "PPL-700", "moderate"),
    (r"\bundead\b|revenants?|draugr|deathless", "PPL-600", "high"),
    (r"\bvampires?\b", "PPL-607", "high"),
    (r"werewolves?|lycanthrop", "STA-201", "high"),
    (r"\belves?\b|elven", "PPL-102", "high"),
    (r"dark elves?", "PPL-102", "moderate"),
    (r"\bdwar(?:f|ves)\b|dwarven", "PPL-103", "high"),
    (r"halflings?|smallfolk|halffoots?|kender", "PPL-104", "moderate"),
    (r"\borcs?\b", "PPL-106", "high"),
    (r"\bgoblins?\b", "PPL-107", "high"),
    (r"\bgiants?\b|giant lineage|\bvarl\b", "PPL-108", "moderate"),
    (r"beastfolk|animal peoples?|talking animals?|hybrid peoples?", "PPL-131", "moderate"),
    (r"insect peoples?", "PPL-118", "high"),
    (r"dragons?|draconic|dragon peoples?", "PPL-800", "moderate"),
    (r"eldritch beings?|cosmic beings?|outer gods?", "PPL-1100", "moderate"),
    (r"monsters?|monster ecology|creature ecology", "PPL-000", "low"),
    (r"peoples?|species|races?|ancestr", "PPL-100", "low"),
    (r"barbarian", "ROL-102", "high"),
    (r"warriors?|soldiers?|martial roles?", "ROL-101", "moderate"),
    (r"knights?|chivalr", "ROL-115", "moderate"),
    (r"paladins?|holy knights?", "ROL-117", "high"),
    (r"samurai", "ROL-120", "high"),
    (r"spearwoman|spearmen?|lancers?", "ROL-110", "moderate"),
    (r"archers?", "ROL-124", "high"),
    (r"rangers?", "ROL-125", "high"),
    (r"hunters?", "ROL-126", "moderate"),
    (r"monster hunters?", "ROL-144", "high"),
    (r"witchers?", "ROL-810", "high"),
    (r"dragon riders?|dragon bonding", "ROL-150", "high"),
    (r"mercenar", "ROL-137", "high"),
    (r"bodyguards?", "ROL-139", "high"),
    (r"pirates?|corsairs?", "ROL-135", "high"),
    (r"thieves?", "ROL-201", "high"),
    (r"rogues?", "ROL-205", "high"),
    (r"assassins?", "ROL-208", "high"),
    (r"con artists?|tricksters?", "ROL-217", "moderate"),
    (r"merchants?|maritime trade", "ROL-702", "moderate"),
    (r"sorcerers?|sorceress", "ROL-302", "high"),
    (r"\bwitches?\b", "ROL-303", "high"),
    (r"wizards?|\bmages?\b|magicians?", "ROL-301", "moderate"),
    (r"necromancers?|necromancer houses?", "ROL-315", "high"),
    (r"summoners?", "ROL-314", "high"),
    (r"shamans?|shamanic", "ROL-502", "high"),
    (r"druids?", "ROL-501", "high"),
    (r"priests?|priestess|clerics?|religious offices?", "ROL-400", "moderate"),
    (r"miko|shrine roles?", "ROL-404", "moderate"),
    (r"onmy[ōo]ji", "ROL-307", "high"),
    (r"diviners?|divination|oracles?", "ROL-316", "moderate"),
    (r"seers?|prophets?", "ROL-409", "moderate"),
    (r"healers?|physicians?", "ROL-601", "moderate"),
    (r"alchemists?|alchemy", "ROL-338", "moderate"),
    (r"smiths?|smithing|weaponsmith", "ROL-610", "moderate"),
    (r"bards?|griots?|storytellers?|chronicler", "ROL-606", "moderate"),
    (r"teachers?|mentors?|masters?", "ROL-628", "moderate"),
    (r"scholars?|apothecar", "ROL-600", "low"),
    (r"kings?|queens?|kingship|queenship|monarch", "ROL-700", "moderate"),
    (r"princes?|princess", "ROL-701", "moderate"),
    (r"courtiers?|court offices?", "ROL-703", "moderate"),
    (r"cultivators?|cultivation", "ROL-906", "high"),
    (r"adventurers?|delvers?|wandering adventurer", "ROL-901", "moderate"),
    (r"ranked hunters?", "ROL-926", "high"),
    (r"inventors?", "ROL-616", "high"),
    (r"craft(?:ing)? roles?|life skills?", "ROL-639", "moderate"),
    (r"elemental (martial )?traditions?|elemental powers?", "PWR-200", "moderate"),
    (r"fire magic|fire origins?", "PWR-201", "moderate"),
    (r"earth power", "PWR-205", "high"),
    (r"blood magic|blood ministration", "PWR-305", "moderate"),
    (r"necromancy|soul magic|death magic", "PWR-303", "moderate"),
    (r"true names?|naming|name loss", "PWR-503", "high"),
    (r"song magic|music magic|guqin", "PWR-506", "moderate"),
    (r"runes?|rune magic", "PWR-504", "moderate"),
    (r"talismans?|seals?", "PWR-511", "moderate"),
    (r"grimoires?", "ART-401", "moderate"),
    (r"psionics?", "PWR-402", "high"),
    (r"ki\b|qi\b|internal refinement", "PWR-115", "moderate"),
    (r"pacts?|contracts?|bargains?", "PWR-104", "moderate"),
    (r"divine powers?|divine favor|god boons?", "PWR-103", "moderate"),
    (r"prophecy|prophetic", "PWR-407", "high"),
    (r"possession|possessed", "PWR-112", "high"),
    (r"mutations?|alchemical enhancement", "STA-206", "moderate"),
    (r"corruption|taint", "STA-421", "moderate"),
    (r"transformation|shapechange|evolutions?|awakenings?", "STA-200", "low"),
    (r"beast transformation", "STA-211", "high"),
    (r"demon transformation", "STA-208", "high"),
    (r"dragon transformation", "STA-210", "high"),
    (r"reincarnat", "STA-501", "high"),
    (r"resurrection|return-by-death", "STA-113", "moderate"),
    (r"immortality|immortal protagonist", "PWR-312", "moderate"),
    (r"soul bond|soul bonding", "STA-307", "high"),
    (r"dragon bond", "STA-409", "high"),
    (r"spirit contracts?|familiar bond|familiar", "STA-308", "moderate"),
    (r"curses?|cursed", "STA-311", "moderate"),
    (r"exile", "STA-504", "moderate"),
    (r"schools?|academ(?:y|ies)|university magic", "INS-309", "moderate"),
    (r"guilds?|adventurer guild", "INS-100", "moderate"),
    (r"thieves['’]? guild", "INS-103", "high"),
    (r"orders?|corps|squads?", "INS-200", "low"),
    (r"knightly orders?|round table", "INS-204", "moderate"),
    (r"monster[- ]hunter order|slayer corps", "INS-208", "moderate"),
    (r"clans?|houses?|dynast", "INS-400", "moderate"),
    (r"sects?|cultivation clans?", "INS-312", "moderate"),
    (r"factions?|political factions?|city factions?", "INS-000", "low"),
    (r"military units?|war hosts?|armies?", "INS-201", "moderate"),
    (r"secret organi[sz]ation", "INS-414", "moderate"),
    (r"church|temple", "INS-300", "low"),
    (r"hero(?:ic)? (?:role|protagonist|king)|\bhero\b", "NAR-100", "low"),
    (r"chosen (?:children|hero|identity)|chosen one", "NAR-101", "moderate"),
    (r"antihero|anti-hero", "NAR-111", "high"),
    (r"tragic (?:champions?|hero)|doomed", "NAR-112", "moderate"),
    (r"dark lord|demon king|demon lord", "NAR-301", "moderate"),
    (r"companions?|party|fellowship|crew", "NAR-200", "low"),
    (r"artifact destruction|destroy.*artifact", "NAR-402", "high"),
    (r"artifact quest|holy relics?|grail", "NAR-401", "moderate"),
    (r"underworld journeys?|underworld trials?|descent", "NAR-410", "moderate"),
    (r"portal fantasy|isekai", "NAR-411", "moderate"),
    (r"exile and return|returning king", "NAR-108", "moderate"),
    (r"heists?", "NAR-434", "high"),
    (r"tournaments?", "NAR-412", "high"),
    (r"apocalypse|apocalyptic", "NAR-418", "moderate"),
    (r"cyclical ages?|world cycles?|fading ages?", "NAR-420", "moderate"),
    (r"quests?|pilgrimage|epic journey", "NAR-400", "low"),
    (r"empires?|imperial|kingdoms?|nations?|polities?", "CIV-100", "low"),
    (r"cities?|urban|settlements?", "CIV-200", "low"),
    (r"cultures?|societ(?:y|ies)|civilizations?", "CIV-000", "low"),
    (r"caste oppression|class oppression|castes?", "CIV-115", "moderate"),
    (r"colonization|colonial", "CIV-405", "moderate"),
    (r"post-apocalyp|dying world|fallen kingdom", "CIV-402", "moderate"),
    (r"magitech|industrial magic|magic.technology|science[- ]fantasy", "CIV-312", "moderate"),
    (r"sacred weapons?|divine weapons?|astras?", "ART-201", "high"),
    (r"enchanted weapons?|living sword|trick weapons?|shard weapons?", "ART-200", "moderate"),
    (r"relics?|talismans?|legendary objects?", "ART-100", "low"),
    (r"chariots?|magical ships?|vehicles?", "ART-500", "moderate"),
    (r"grimoires?|magical texts?|book", "ART-400", "low"),
    (r"spirit world|divine realms?|parallel worlds?|shadow worlds?|multiverse", "COS-400", "moderate"),
    (r"underworld|afterlife|death realm", "COS-300", "moderate"),
    (r"sacred landscape|world tree|world mountain|sacred geography", "COS-500", "moderate"),
    (r"cosmic time|world ages?|cyclical apocalypse", "COS-700", "moderate"),
    (r"cosmology|cosmologies|living worlds?", "COS-000", "low"),
    (r"vows?|oaths?|geasa", "LAW-100", "moderate"),
    (r"boons?", "LAW-201", "high"),
    (r"curses?", "LAW-204", "high"),
    (r"sacrifice|offerings?|equivalent exchange", "LAW-300", "moderate"),
    (r"purity|pollution", "LAW-400", "moderate"),
    (r"fate|destiny|world tendency", "LAW-501", "moderate"),
    (r"dharma|cosmic order|balance", "LAW-500", "moderate"),
    (r"death and rebirth|reincarnation cycle", "LAW-700", "moderate"),
    (r"kinship|genealogy|lineage|bloodline", "LAW-800", "low"),
    (r"race/class|class/race|classes?|jobs?|vocations?|professions?|playbooks?", "MEC-301", "low"),
    (r"job system|class changes?|promotion|multiclass", "MEC-302", "moderate"),
    (r"skill trees?|classless builds?|builds?", "MEC-300", "low"),
    (r"alignment|morality", "MEC-322", "moderate"),
    (r"permadeath|permanent loss", "MEC-205", "high"),
    (r"stress|affliction|virtue", "MEC-218", "moderate"),
    (r"resource progression|extreme leveling|power progression", "MEC-300", "low"),
    (r"faction.*unit|unit tiers?|army factions?", "MEC-428", "moderate"),
]


SOURCE_RELATIONSHIPS: list[tuple[str, str, str, str]] = [
    ("SRC-026", "contains-work", "SRC-031", "The Lord of the Rings belongs to Tolkien's literary legendarium."),
    ("SRC-031", "adapted-as", "SRC-266", "The film trilogy adapts the novel."),
    ("SRC-026", "adapted-as", "SRC-267", "Screen continuity draws on Tolkien's Second Age material under separate adaptation constraints."),
    ("SRC-046", "adapted-as", "SRC-264", "Television adaptation of the novel sequence, with material divergences."),
    ("SRC-046", "adapted-as", "SRC-265", "Television prequel continuity draws primarily on Fire & Blood material within the same world."),
    ("SRC-047", "adapted-as", "SRC-161", "Game continuity adapts and extends the literary Witcher world."),
    ("SRC-047", "adapted-as", "SRC-268", "Television continuity adapts the literary Witcher material."),
    ("SRC-044", "adapted-as", "SRC-269", "Television adaptation of The Wheel of Time."),
    ("SRC-247", "prequel-continuity", "SRC-276", "Age of Resistance is a prequel series in The Dark Crystal screen world."),
    ("SRC-027", "adapted-as", "SRC-253", "Film traditions adapt and consolidate the Conan literary icon."),
    ("SRC-005", "retold-as", "SRC-033", "Modern Arthurian literary retelling."),
    ("SRC-005", "adapted-as", "SRC-254", "Cinematic Arthurian adaptation."),
    ("SRC-035", "same-corpus-as", "SRC-085", "Earthsea and The Books of Earthsea overlap as source-corpus labels and require deduplication."),
    ("SRC-043", "same-setting-as", "SRC-083", "Both records refer to Osten Ard material at different scope."),
    ("SRC-014", "witnessed-by", "SRC-277", "The BORI Critical Edition is a reference witness for the wider epic tradition."),
    ("SRC-014", "contains-text", "SRC-278", "The Bhagavad Gītā is embedded in the Mahābhārata."),
    ("SRC-277", "contains-text", "SRC-278", "Critical-edition indexing situates the Gītā within Bhīṣma Parva."),
    ("SRC-018", "represented-by-work", "SRC-077", "The novel is a major modern xianxia/cultivation example."),
    ("SRC-018", "represented-by-work", "SRC-078", "The novel is a cultivation-fiction example with distinct clan and demonic-cultivation emphases."),
    ("SRC-018", "represented-by-work", "SRC-240", "Soul Land mechanizes cultivation through martial souls and spirit rings."),
    ("SRC-016", "draws-from", "SRC-279", "Journey to the West participates in overlapping Chinese religious and literary systems."),
    ("SRC-017", "draws-from", "SRC-279", "Investiture of the Gods adapts and systematizes Chinese divine and immortal traditions."),
    ("SRC-089", "setting", "SRC-090", "Forgotten Realms is a D&D setting."),
    ("SRC-089", "setting", "SRC-091", "Greyhawk is a D&D setting."),
    ("SRC-089", "setting", "SRC-092", "Dragonlance is a D&D setting and literary line."),
    ("SRC-089", "setting", "SRC-093", "Eberron is a D&D setting."),
    ("SRC-089", "setting", "SRC-094", "Dark Sun is a D&D setting."),
    ("SRC-089", "setting", "SRC-095", "Planescape is a D&D setting."),
    ("SRC-089", "setting", "SRC-096", "Ravenloft is a D&D setting."),
    ("SRC-089", "setting", "SRC-097", "Spelljammer is a D&D setting."),
    ("SRC-089", "implemented-as", "SRC-139", "Baldur's Gate implements D&D rules in a CRPG series."),
    ("SRC-089", "implemented-as", "SRC-140", "Icewind Dale implements D&D rules in a CRPG."),
    ("SRC-095", "implemented-as", "SRC-141", "Planescape: Torment is a game set in Planescape."),
    ("SRC-089", "implemented-as", "SRC-142", "Neverwinter Nights implements D&D rules and setting material."),
    ("SRC-098", "implemented-as", "SRC-146", "Kingmaker implements Pathfinder rules."),
    ("SRC-098", "implemented-as", "SRC-147", "Wrath of the Righteous implements Pathfinder rules."),
    ("SRC-107", "implemented-as", "SRC-192", "The Returns trilogy adapts the Shadowrun tabletop setting."),
    ("SRC-127", "versioned-as", "SRC-128", "Reborn is a version-specific Tactics Ogre record."),
    ("SRC-129", "franchise-entry", "SRC-130", "Final Fantasy Tactics is a franchise branch with a distinct tactical job system."),
    ("SRC-129", "franchise-entry", "SRC-131", "Final Fantasy XIV is an MMO entry with its own continuity and job snapshots."),
    ("SRC-135", "franchise-entry", "SRC-136", "World of Warcraft is the MMO branch of Warcraft."),
    ("SRC-001", "adapted-by", "SRC-193", "Hades mechanizes Greek deities, boons, underworld offices, and repeated escape."),
    ("SRC-001", "adapted-by", "SRC-274", "The series freely adapts Greek heroic myth."),
    ("SRC-001", "adapted-by", "SRC-273", "The series syncretically adapts Greek myth among other traditions."),
    ("SRC-001", "adapted-by", "SRC-281", "Age of Mythology factionalizes and mechanizes Greek mythic material."),
    ("SRC-002", "adapted-by", "SRC-281", "Age of Mythology factionalizes and mechanizes Norse mythic material."),
    ("SRC-013", "adapted-by", "SRC-281", "Age of Mythology factionalizes and mechanizes Egyptian mythic material."),
    ("SRC-279", "adapted-by", "SRC-281", "Later game expansions adapt selected Chinese mythic material."),
    ("SRC-019", "adapted-by", "SRC-202", "Inuyasha adapts and recombines yōkai and shrine traditions."),
    ("SRC-019", "adapted-by", "SRC-231", "Spirited Away draws on kami/yōkai imagery while creating a distinct fictional spirit economy."),
    ("SRC-019", "adapted-by", "SRC-233", "Jujutsu Kaisen reworks curse and exorcist traditions into a modern battle system."),
    ("SRC-005", "influences", "SRC-255", "Dragonheart reworks knight-and-dragon materials through a soul-bond premise."),
    ("SRC-089", "adapted-as", "SRC-271", "Vox Machina adapts a D&D actual-play campaign and its class vocabulary."),
    ("SRC-089", "adapted-as", "SRC-272", "Honor Among Thieves adapts D&D classes, spells, monsters, and setting conventions to film."),
]


def clean(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def sheet_records(sheet: Any, header_row: int = 1) -> list[dict[str, Any]]:
    headers = [clean(cell.value) for cell in sheet[header_row]]
    output = []
    for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        item = {header: value for header, value in zip(headers, row) if header and value not in (None, "")}
        if item:
            output.append(item)
    return output


def split_contribution(value: str) -> list[str]:
    value = value.strip().rstrip(".")
    parts = [part.strip() for part in re.split(r"[;,]", value) if part.strip()]
    return parts or [value]


def map_concept(concept: str) -> list[tuple[str, str]]:
    matches: list[tuple[str, str]] = []
    lowered = concept.casefold()
    for pattern, archetype_id, confidence in CONCEPT_RULES:
        if re.search(pattern, lowered, re.IGNORECASE) and archetype_id not in {item[0] for item in matches}:
            matches.append((archetype_id, confidence))
        if len(matches) >= 3:
            break
    return matches


def add_table(sheet: Any, display_name: str) -> None:
    table = Table(displayName=display_name, ref=f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False
    )
    sheet.add_table(table)


def style_sheet(sheet: Any, widths: list[int], freeze: str = "A2") -> None:
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = freeze
    sheet.auto_filter.ref = sheet.dimensions
    side = Side(style="thin", color=GRID)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=HEADER)
        cell.font = Font(color=INK, bold=True, size=10)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(bottom=side)
    sheet.row_dimensions[1].height = 32
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(color="26344A", size=9)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=Side(style="hair", color="D8E0EC"))
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--recalculate", action="store_true")
    args = parser.parse_args()
    fingerprint = source_fingerprint([args.input, Path(__file__)])

    workbook = load_workbook(args.input)
    for name in ["Source Concept Map", "Source Relationships", "Source Coverage Audit", "Visualization Guide"]:
        if name in workbook.sheetnames:
            del workbook[name]

    taxonomy = sheet_records(workbook["Master Taxonomy"])
    sources = sheet_records(workbook["Source Corpus"])
    entries = sheet_records(workbook["Seed Catalogue"])
    priorities = sheet_records(workbook["Source Priority"], 6)
    archetype_names = {clean(row["Archetype_ID"]): clean(row["Preferred_Name"]) for row in taxonomy}
    source_names = {clean(row["Source_ID"]): clean(row["Title_or_Franchise"]) for row in sources}
    entry_counts = Counter(clean(row.get("Source_ID")) for row in entries)
    priority_by_source = {clean(row.get("Source_ID")): row for row in priorities}

    missing_rule_ids = sorted({archetype_id for _, archetype_id, _ in CONCEPT_RULES if archetype_id not in archetype_names})
    if missing_rule_ids:
        raise SystemExit(f"Concept rules reference unknown archetypes: {missing_rule_ids}")

    concept_sheet = workbook.create_sheet("Source Concept Map")
    concept_headers = [
        "Concept_ID",
        "Source_ID",
        "Source_Title",
        "Orientation_Concept",
        "Archetype_IDs",
        "Archetype_Names",
        "Relationship",
        "Mapping_Basis",
        "Confidence",
        "Evidence_Level",
        "Review_Status",
        "Reference_URL",
        "Notes",
    ]
    concept_sheet.append(concept_headers)
    concept_counts: Counter[str] = Counter()
    normalized_counts: Counter[str] = Counter()
    concept_index = 1
    confidence_order = {"high": 3, "moderate": 2, "low": 1}
    for source in sources:
        source_id = clean(source["Source_ID"])
        concepts = split_contribution(clean(source.get("Distinctive_Contribution")))
        for concept in concepts:
            mappings = map_concept(concept)
            ids = [item[0] for item in mappings]
            confidences = [item[1] for item in mappings]
            confidence = min(confidences, key=lambda value: confidence_order[value]) if confidences else "unmapped"
            concept_sheet.append(
                [
                    f"CON-{concept_index:04d}",
                    source_id,
                    source_names[source_id],
                    concept,
                    "; ".join(ids),
                    "; ".join(archetype_names[item] for item in ids),
                    "orientation-maps-to" if ids else "source-describes-concept",
                    "Source Corpus.Distinctive_Contribution",
                    confidence,
                    "Corpus orientation only",
                    "Needs source-specific verification",
                    clean(source.get("Reference_URL")),
                    "Comparative navigation coordinate; not a canonical source entry or claim of exact equivalence.",
                ]
            )
            concept_counts[source_id] += 1
            if ids:
                normalized_counts[source_id] += 1
            concept_index += 1
    style_sheet(concept_sheet, [14, 12, 30, 34, 22, 34, 22, 31, 12, 22, 27, 35, 48])
    add_table(concept_sheet, "SourceConceptMapTable")

    relationship_sheet = workbook.create_sheet("Source Relationships")
    relationship_sheet.append(
        [
            "Relationship_ID",
            "From_Source_ID",
            "From_Title",
            "Relationship",
            "To_Source_ID",
            "To_Title",
            "Scope_Note",
            "Evidence_Level",
            "Review_Status",
        ]
    )
    for index, (from_id, relation, to_id, note) in enumerate(SOURCE_RELATIONSHIPS, 1):
        if from_id not in source_names or to_id not in source_names:
            raise SystemExit(f"Source relationship references unknown source: {from_id} -> {to_id}")
        relationship_sheet.append(
            [
                f"REL-{index:04d}",
                from_id,
                source_names[from_id],
                relation,
                to_id,
                source_names[to_id],
                note,
                "Corpus-level relationship reviewed",
                "Editorial relationship — verify during deep pass",
            ]
        )
    style_sheet(relationship_sheet, [16, 15, 32, 22, 15, 32, 65, 31, 38])
    add_table(relationship_sheet, "SourceRelationshipsTable")

    coverage_sheet = workbook.create_sheet("Source Coverage Audit")
    coverage_sheet.append(
        [
            "Source_ID",
            "Title_or_Franchise",
            "Medium",
            "Region_Tradition",
            "Priority_Tier",
            "Baseline_Rank",
            "Recommended_Treatment",
            "Canonical_Entry_Count",
            "Orientation_Concept_Count",
            "Normalized_Concept_Count",
            "Canonical_Target",
            "Canonical_Coverage_Status",
            "Orientation_Coverage_Status",
            "Next_Action",
        ]
    )
    status_counts: Counter[str] = Counter()
    for source in sources:
        source_id = clean(source["Source_ID"])
        priority = priority_by_source.get(source_id, {})
        treatment = clean(priority.get("Recommended_Treatment")) or "Representative pass"
        target = target_for_treatment(treatment)
        canonical_count = entry_counts[source_id]
        if canonical_count == 0:
            canonical_status = "Orientation only"
        elif canonical_count < target * 0.4:
            canonical_status = "Canonical scouting"
        elif canonical_count < target:
            canonical_status = "Canonical mapping in progress"
        else:
            canonical_status = "Canonical pass target reached"
        status_counts[canonical_status] += 1
        coverage_sheet.append(
            [
                source_id,
                source_names[source_id],
                clean(source.get("Medium")),
                clean(source.get("Region_Tradition")),
                clean(source.get("Priority_Tier")),
                priority.get("Baseline_Rank", ""),
                treatment,
                canonical_count,
                concept_counts[source_id],
                normalized_counts[source_id],
                target,
                canonical_status,
                "Orientation mapped" if concept_counts[source_id] else "Missing orientation",
                "Verify canonical terms and citations" if canonical_count == 0 else "Extend or audit canonical source pass",
            ]
        )
    style_sheet(coverage_sheet, [12, 32, 28, 30, 14, 14, 25, 18, 21, 22, 17, 29, 25, 38])
    add_table(coverage_sheet, "SourceCoverageAuditTable")
    coverage_sheet.conditional_formatting.add(
        f"H2:J{coverage_sheet.max_row}",
        ColorScaleRule(start_type="min", start_color="FEE2E2", mid_type="percentile", mid_value=50, mid_color="FEF3C7", end_type="max", end_color="D1FAE5"),
    )

    summary_start = coverage_sheet.max_row + 3
    coverage_sheet.cell(summary_start, 1, "Canonical coverage status")
    coverage_sheet.cell(summary_start, 2, "Sources")
    for offset, (status, count) in enumerate(status_counts.items(), 1):
        coverage_sheet.cell(summary_start + offset, 1, status)
        coverage_sheet.cell(summary_start + offset, 2, count)
    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = "Canonical source-pass coverage"
    chart.y_axis.title = "Status"
    chart.x_axis.title = "Sources"
    data = Reference(coverage_sheet, min_col=2, min_row=summary_start, max_row=summary_start + len(status_counts))
    categories = Reference(coverage_sheet, min_col=1, min_row=summary_start + 1, max_row=summary_start + len(status_counts))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.height = 7
    chart.width = 13
    coverage_sheet.add_chart(chart, f"D{summary_start}")

    guide = workbook.create_sheet("Visualization Guide", 1)
    guide.sheet_view.showGridLines = False
    guide.column_dimensions["A"].width = 3
    guide.column_dimensions["B"].width = 28
    guide.column_dimensions["C"].width = 92
    guide.merge_cells("B2:C2")
    guide["B2"] = "Version 3 Interactive Relationship Atlas"
    guide["B2"].font = Font(name="Georgia", size=24, bold=True, color=INK)
    guide["B2"].fill = PatternFill("solid", fgColor=NAVY)
    guide["B3"] = "Deliverable"
    guide["C3"] = "TypeScript + Sigma.js + Graphology WebGL explorer generated from this workbook."
    guide["B4"] = "Run locally"
    guide["C4"] = "From the project directory: npm install, then npm run dev. Open the local URL printed by Vite."
    guide["B5"] = "Build"
    guide["C5"] = "npm run build regenerates this workbook's JSON graph projection, type-checks the app, and creates dist/."
    guide["B7"] = "Semantic zoom"
    guide["C7"] = "Far: eleven domains and all sources. Middle: archetype families and source bridges. Near: source concepts, canonical entries, adaptations, evidence, and citations."
    guide["B8"] = "Coverage view"
    guide["C8"] = "Separates orientation coverage (281/281 sources) from canonical source-pass coverage (currently 21/281 sources)."
    guide["B9"] = "Integrity rule"
    guide["C9"] = "Source Concept Map rows are orientation coordinates derived from editorial corpus notes. They must not be cited as canonical claims. Seed Catalogue remains the canonical source-entry layer."
    guide["B11"] = "Graph node types"
    guide["C11"] = "Domain · normalized archetype · source/tradition · orientation concept · canonical source entry · adaptation crosswalk."
    guide["B12"] = "Graph edge types"
    guide["C12"] = "taxonomy · canonical mapping · orientation mapping · source containment · adaptation · version · influence · implementation."
    guide["B14"] = "Version 3 counts"
    guide["C14"] = f"{len(sources)} sources · {len(taxonomy)} archetypes · {len(entries)} canonical entries · {concept_index - 1} orientation concepts · {len(SOURCE_RELATIONSHIPS)} source relationships."
    for row in range(2, 15):
        guide.cell(row, 2).alignment = Alignment(vertical="top", wrap_text=True)
        guide.cell(row, 3).alignment = Alignment(vertical="top", wrap_text=True)
        if row >= 3:
            guide.cell(row, 2).font = Font(bold=True, color="394A65", size=10)
            guide.cell(row, 3).font = Font(color="52627A", size=10)
        guide.row_dimensions[row].height = 32 if row != 2 else 44
    guide.freeze_panes = "B3"

    start = workbook["Start Here"]
    start["B5"] = "3.0"
    start["A6"] = "Source fingerprint"
    start["B6"] = fingerprint
    start["A3"] = "A source-faithful comparative ontology with complete corpus orientation coverage and a WebGL relationship explorer"

    dictionary = workbook["Data Dictionary"]
    new_dictionary_rows = [
        ("Source Concept Map", "Concept_ID", "Text / key", "Yes", "Stable orientation-concept identifier.", "Not a canonical source-entry ID."),
        ("Source Concept Map", "Orientation_Concept", "Text", "Yes", "Contribution phrase from Source Corpus.", "Requires source-specific verification before promotion."),
        ("Source Concept Map", "Archetype_IDs", "Delimited foreign keys", "No", "One to three normalized comparison coordinates.", "Partial analogy only unless separately reviewed."),
        ("Source Relationships", "Relationship", "Controlled text", "Yes", "Typed relation between two source records.", "Direction is From Source to To Source."),
        ("Source Coverage Audit", "Canonical_Coverage_Status", "Controlled text", "Yes", "Progress against the tiered source-pass target.", "Never infer completion from corpus inclusion."),
        ("Source Coverage Audit", "Orientation_Coverage_Status", "Controlled text", "Yes", "Whether the source has navigation-level concept records.", "Orientation is not canonical verification."),
    ]
    for row in new_dictionary_rows:
        dictionary.append(row)

    workbook.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True, calcOnSave=True)
    workbook.save(args.output)

    office_binary = shutil.which("libreoffice") or shutil.which("soffice")
    if args.recalculate and office_binary:
        with tempfile.TemporaryDirectory(prefix="fantasy-atlas-recalc-") as temporary_directory:
            subprocess.run(
                [
                    office_binary,
                    "--headless",
                    "--convert-to",
                    "xlsx",
                    "--outdir",
                    temporary_directory,
                    str(args.output),
                ],
                check=True,
                capture_output=True,
                text=True,
            )
            recalculated = Path(temporary_directory) / args.output.name
            if not recalculated.exists():
                raise SystemExit("LibreOffice completed without producing the recalculated workbook")
            shutil.copy2(recalculated, args.output)
    elif args.recalculate:
        raise SystemExit("LibreOffice was requested but is not available")

    normalize_xlsx(args.output)

    if set(source_names) != set(concept_counts):
        missing = sorted(set(source_names) - set(concept_counts))
        raise SystemExit(f"Orientation coverage incomplete; missing sources: {missing}")

    print(
        f"Wrote {args.output.name}: {len(sources)} / {len(sources)} sources orientation-mapped, "
        f"{sum(1 for count in entry_counts.values() if count)} canonically seeded, "
        f"{concept_index - 1} source concepts ({sum(normalized_counts.values())} normalized), "
        f"{len(SOURCE_RELATIONSHIPS)} source relationships."
    )


if __name__ == "__main__":
    main()
