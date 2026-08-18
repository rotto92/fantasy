#!/usr/bin/env python3
"""Build the character-first Version 4 workbook from Version 3 and validated research."""

from __future__ import annotations

import argparse
import json
import shutil
import subprocess
import tempfile
from collections import Counter
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.table import Table, TableStyleInfo

from reproducible import normalize_xlsx, source_fingerprint


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = ROOT / "fantasy_high_fantasy_archetype_atlas_v3.xlsx"
DEFAULT_RESEARCH = ROOT / "public" / "data" / "characters.json"
DEFAULT_OUTPUT = ROOT / "fantasy_high_fantasy_archetype_atlas_v4.xlsx"

NAVY = "0B1222"
HEADER = "17243B"
INK = "EAF1FB"
MUTED = "A7B5C9"
GOLD = "EFC87A"
TEAL = "70D6C8"
VIOLET = "A99AF6"
GRID = "2B3A55"


def joined(values: Iterable[Any], separator: str = "; ") -> str:
    return separator.join(str(value) for value in values if value not in (None, ""))


def citation_summary(citations: list[dict[str, Any]]) -> str:
    return " | ".join(
        f"{citation.get('locator', '')}: {joined(citation.get('supports', []), '; ')} [{citation.get('url', '')}]"
        for citation in citations
    )


def add_sheet(workbook: Any, name: str, headers: list[str], rows: list[list[Any]]) -> Any:
    if name in workbook.sheetnames:
        del workbook[name]
    sheet = workbook.create_sheet(name)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{sheet.cell(1, len(headers)).coordinate}"
    thin = Side(style="thin", color=GRID)
    for column, header in enumerate(headers, 1):
        cell = sheet.cell(1, column, header)
        cell.fill = PatternFill("solid", fgColor=HEADER)
        cell.font = Font(color=INK, bold=True, size=10)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)
    sheet.row_dimensions[1].height = 32
    for row_index, row in enumerate(rows, 2):
        for column, value in enumerate(row, 1):
            cell = sheet.cell(row_index, column, value)
            cell.font = Font(color=INK if column <= 3 else MUTED, size=9)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=Side(style="hair", color=GRID))
            if row_index % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="0E1829")
            if "URL" in headers[column - 1] and isinstance(value, str) and value.startswith(("http://", "https://")):
                cell.hyperlink = value
                cell.font = Font(color=TEAL, underline="single", size=9)
        sheet.row_dimensions[row_index].height = 31
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.sheet_format.defaultColWidth = 14
    if rows:
        table = Table(displayName=f"tbl{name.replace(' ', '')}", ref=f"A1:{sheet.cell(len(rows) + 1, len(headers)).coordinate}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False, showLastColumn=False)
        sheet.add_table(table)
    return sheet


def set_widths(sheet: Any, widths: dict[str, float]) -> None:
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--research", type=Path, default=DEFAULT_RESEARCH)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--recalculate", action="store_true")
    args = parser.parse_args()
    fingerprint = source_fingerprint([args.input, args.research, Path(__file__)])

    if not args.research.exists():
        raise SystemExit(f"Validated character payload not found: {args.research}")
    data = json.loads(args.research.read_text(encoding="utf-8"))
    shutil.copy2(args.input, args.output)
    workbook = load_workbook(args.output)
    if "Visualization Guide" in workbook.sheetnames:
        workbook["Visualization Guide"].title = "V3 Viz Archive"

    characters = data["characters"]
    relationships = data["relationships"]
    terms = data["sourceTerms"]
    audits = data["sources"]
    corpus = data["corpusSources"]
    taxonomy = data["taxonomy"]
    review_index = data.get("independentReviews", {})
    source_review_status = review_index.get("source_review_status", {})
    audit_by_source = {row["source_id"]: row for row in audits}
    character_counts = Counter(row["source_id"] for row in characters)
    relationship_counts = Counter(row["source_id"] for row in relationships)
    term_counts = Counter(row["source_id"] for row in terms)

    character_headers = [
        "Character_ID", "Canonical_Name", "Aliases", "Source_ID", "Source_Title", "Continuity",
        "Work_or_Witness", "Character_Kind", "Description", "Canon_Status", "Evidence_Level",
        "Citation_Count", "Primary_Reference_URL", "Citation_Summary", "Comparison_Cautions", "Spoiler_Level", "Review_Status",
        "Independent_Second_Review",
    ]
    character_rows = [
        [
            row["character_id"], row["canonical_name"], joined(row.get("aliases", [])), row["source_id"],
            row["source_title"], row["continuity"], row["work_or_witness"], row["character_kind"],
            row["description"], row["canon_status"], row["evidence_level"], len(row.get("citations", [])),
            row.get("citations", [{}])[0].get("url", "") if row.get("citations") else "",
            citation_summary(row.get("citations", [])), joined(row.get("comparison_cautions", [])),
            row["spoiler_level"], row["review_status"],
            "source-focused review recorded; individual claim review remains scoped"
            if row["source_id"] in source_review_status
            else "pending",
        ]
        for row in characters
    ]
    sheet = add_sheet(workbook, "Character Catalogue", character_headers, character_rows)
    set_widths(sheet, {"A": 19, "B": 24, "C": 22, "D": 12, "E": 29, "F": 31, "G": 31, "H": 15, "I": 50, "J": 18, "K": 18, "L": 13, "M": 34, "N": 70, "O": 48, "P": 13, "Q": 15, "R": 20})

    dimension_headers = [
        "Dimension_Link_ID", "Character_ID", "Character_Name", "Source_ID", "Dimension",
        "Source_Native_Term", "Archetype_IDs", "Archetype_Names", "Confidence", "Mapping_Note",
    ]
    dimension_rows: list[list[Any]] = []
    link_index = 1
    for character in characters:
        for dimension, values in character["dimensions"].items():
            for value in values:
                archetype_ids = value.get("archetype_ids", [])
                dimension_rows.append(
                    [
                        f"CDM-{link_index:05d}", character["character_id"], character["canonical_name"],
                        character["source_id"], dimension, value.get("term", ""), joined(archetype_ids),
                        joined(taxonomy.get(archetype_id, {}).get("name", archetype_id) for archetype_id in archetype_ids),
                        value.get("confidence", ""), value.get("note", ""),
                    ]
                )
                link_index += 1
    sheet = add_sheet(workbook, "Character Dimensions", dimension_headers, dimension_rows)
    set_widths(sheet, {"A": 18, "B": 19, "C": 24, "D": 12, "E": 30, "F": 30, "G": 22, "H": 32, "I": 13, "J": 52})

    relationship_headers = [
        "Relationship_ID", "Source_ID", "Source_Character_ID", "Source_Character_Name",
        "Target_Character_ID", "Target_Character_Name", "Relationship_Type", "Label", "Direction",
        "Continuity", "Confidence", "Note", "Primary_Reference_URL", "Citation_Summary",
    ]
    character_by_id = {row["character_id"]: row for row in characters}
    relationship_rows = [
        [
            row["relationship_id"], row["source_id"], row["source_character_id"],
            character_by_id[row["source_character_id"]]["canonical_name"], row["target_character_id"],
            character_by_id[row["target_character_id"]]["canonical_name"], row["relationship_type"],
            row["label"], row["direction"], row["continuity"], row["confidence"], row.get("note", ""),
            row.get("citations", [{}])[0].get("url", "") if row.get("citations") else "",
            citation_summary(row.get("citations", [])),
        ]
        for row in relationships
    ]
    sheet = add_sheet(workbook, "Character Relationships", relationship_headers, relationship_rows)
    set_widths(sheet, {"A": 20, "B": 12, "C": 20, "D": 24, "E": 20, "F": 24, "G": 18, "H": 33, "I": 12, "J": 30, "K": 12, "L": 35, "M": 34, "N": 68})

    term_headers = [
        "Term_ID", "Source_ID", "Work_or_Witness", "Canonical_Term", "Identity_Forms", "Original_Language", "Original_Script",
        "Transliteration", "Literal_Gloss", "Dimension", "Archetype_IDs", "Archetype_Names",
        "Mapping_Relation", "Definition", "Cultural_Caution", "Primary_Reference_URL", "Citation_Summary", "Review_Status",
    ]
    term_rows = []
    for row in terms:
        archetype_ids = row.get("archetype_ids", [])
        term_rows.append(
            [
                row["term_id"], row["source_id"], row["work_or_witness"], row["canonical_term"], joined(row.get("identity_forms", [])), row.get("original_language", ""),
                row.get("original_script", ""), row.get("transliteration", ""), row.get("literal_gloss", ""),
                row["dimension"], joined(archetype_ids),
                joined(taxonomy.get(archetype_id, {}).get("name", archetype_id) for archetype_id in archetype_ids),
                row["mapping_relation"], row["definition"], row["cultural_caution"],
                row.get("citations", [{}])[0].get("url", "") if row.get("citations") else "",
                citation_summary(row.get("citations", [])), row["review_status"],
            ]
        )
    sheet = add_sheet(workbook, "Character Source Terms", term_headers, term_rows)
    set_widths(sheet, {"A": 18, "B": 12, "C": 38, "D": 25, "E": 24, "F": 18, "G": 20, "H": 18, "I": 26, "J": 29, "K": 20, "L": 28, "M": 17, "N": 44, "O": 48, "P": 34, "Q": 68, "R": 15})

    audit_headers = [
        "Source_ID", "Source_Title", "Medium", "Region_Tradition", "Priority_Tier", "Character_Pass_Status",
        "Continuity_Scope", "Coverage_Rule", "In_Scope_Characters", "Completed_Characters",
        "Integrated_Character_Count", "Relationship_Count", "Source_Term_Count", "Omissions", "Uncertainties",
        "Evidence_Basis", "Independent_Second_Review", "Last_Reviewed", "Corpus_Reference_URL",
    ]
    audit_rows = []
    for source in corpus:
        source_id = source["sourceId"]
        audit = audit_by_source.get(source_id)
        audit_rows.append(
            [
                source_id, source["title"], source["medium"], source["region"], source["priorityTier"],
                audit.get("completion_status", "not-started") if audit else "not-started",
                audit.get("continuity_scope", "") if audit else "",
                audit.get("coverage_rule", "") if audit else "",
                audit.get("in_scope_character_count", 0) if audit else 0,
                audit.get("completed_character_count", 0) if audit else 0,
                character_counts[source_id], relationship_counts[source_id], term_counts[source_id],
                joined(audit.get("omissions", [])) if audit else "",
                joined(audit.get("uncertainties", [])) if audit else "",
                audit.get("evidence_basis", "") if audit else "",
                (
                    "focused review recorded: "
                    + joined(source_review_status[source_id].get("review_types", []))
                )
                if source_id in source_review_status
                else ("pending" if audit else "not-applicable"),
                audit.get("last_reviewed", "") if audit else "",
                source["referenceUrl"],
            ]
        )
    sheet = add_sheet(workbook, "Character Research Audit", audit_headers, audit_rows)
    set_widths(sheet, {"A": 12, "B": 31, "C": 24, "D": 25, "E": 13, "F": 22, "G": 48, "H": 54, "I": 14, "J": 14, "K": 15, "L": 13, "M": 13, "N": 55, "O": 55, "P": 50, "Q": 20, "R": 14, "S": 34})

    review_headers = [
        "Source_ID", "Source_Title", "Review_Status", "Review_Types", "Outcomes",
        "Unresolved_Limits", "Ledger_Files",
    ]
    review_rows = []
    for source in corpus:
        source_id = source["sourceId"]
        review = source_review_status.get(source_id)
        review_rows.append(
            [
                source_id,
                source["title"],
                "focused review recorded" if review else "full second review pending",
                joined(review.get("review_types", [])) if review else "",
                joined(review.get("outcomes", []), " | ") if review else "",
                joined(review.get("unresolved_limits", []), " | ") if review else "Focused independent second review remains pending.",
                joined(review.get("ledgers", [])) if review else "",
            ]
        )
    for index, review in enumerate(review_index.get("corpus_wide_reviews", []), 1):
        review_rows.append(
            [
                f"CORPUS-{index:02d}",
                "All compiled source passes",
                "corpus-wide contract review",
                review.get("review_type", ""),
                review.get("outcome", ""),
                joined(review.get("unresolved_limits", []), " | "),
                review.get("ledger", ""),
            ]
        )
    sheet = add_sheet(workbook, "Independent Reviews", review_headers, review_rows)
    set_widths(sheet, {"A": 14, "B": 32, "C": 27, "D": 32, "E": 58, "F": 64, "G": 42})

    guide_headers = ["Principle", "Version_4_Implementation", "Reason"]
    guide_rows = [
        ["One visual noun", "Every visible point is a normalized being/entity or role/vocation concept.", "Keeps source-native characters, sources, and terms as evidence rather than ontologically equivalent graph nodes."],
        ["Evidence is attached", "Characters and source-native terms supply cited examples, dimensions, and provenance for normalized concepts.", "Preserves source identity while allowing concepts to be compared without promoting evidence records into the graph."],
        ["No global edges", "Relationships appear only in a selected normalized concept's local constellation.", "Maintains readable labels and makes every displayed relationship interpretable."],
        ["Families organize concepts", "Tier-2 normalized families group specific archetypes and support bounded zoom and progressive disclosure.", "Provides overview and drill-down without fake source or character nodes."],
        ["Four shipped views", "Constellations maps normalized concepts; Catalogue lists them by family; Relations shows one selected concept's bounded neighborhood; Research reports source-pass and review status.", "Keeps every documented control aligned with the public D3 application while preserving one evidence model across views."],
        ["Coverage is separate", "All corpus sources appear in Character Research Audit and the Research dashboard; focused review is tracked in Independent Reviews.", "A bounded source pass is not confused with exhaustive franchise research or a full second review."],
        ["Evidence first", "Characters, relationships, source terms, and normalized mappings require claim-level citations, locators, and review status.", "Separates accepted evidence from retained or quarantined research metadata."],
    ]
    sheet = add_sheet(workbook, "Character Viz Guide", guide_headers, guide_rows)
    set_widths(sheet, {"A": 25, "B": 70, "C": 70})

    if "Data Dictionary" in workbook.sheetnames:
        dictionary = workbook["Data Dictionary"]
        dictionary_rows = [
            ("Character Catalogue", "Character_ID", "Text / primary key", "Yes", "Stable named-character identifier within a source pass.", "Characters remain evidence records and never become public graph nodes."),
            ("Character Dimensions", "Character_ID", "Foreign key", "Yes", "Links a character to source-native dimensional attributes.", "Archetype mappings may be empty when analogy would distort."),
            ("Character Relationships", "Source_Character_ID / Target_Character_ID", "Foreign keys", "Yes", "Evidence-backed character-to-character relationship.", "Non-character endpoints are prohibited."),
            ("Character Source Terms", "Canonical_Term", "Text", "Yes", "Source-native comparison and filter vocabulary.", "Does not replace linguistic, ritual, or cultural meaning."),
            ("Character Source Terms", "Work_or_Witness", "Text", "Yes", "Claim-specific work, edition, episode, or other bounded witness.", "Source-wide audit scope remains separate."),
            ("Character Research Audit", "Character_Pass_Status", "Controlled text", "Yes", "Completeness against a declared witness scope and coverage rule.", "Independent second review is tracked separately."),
            ("Independent Reviews", "Review_Status", "Controlled text", "Yes", "Focused source-review and corpus-wide contract-audit status.", "Pending means no focused independent second review is yet recorded."),
            ("Character Viz Guide", "Principle", "Text", "Yes", "Version 4 normalized-concept interaction and visual-grammar rule.", "Defines the public explorer's normalized-concept graph and evidence boundary."),
        ]
        for row in dictionary_rows:
            dictionary.append(row)
        for table in dictionary.tables.values():
            table.ref = f"A1:F{dictionary.max_row}"

    if "Start Here" in workbook.sheetnames:
        start = workbook["Start Here"]
        start["A1"] = "Fantasy & High-Fantasy Comparative Archetype Atlas — Version 4.0"
        start["A1"].font = Font(color=GOLD, bold=True, size=18)
        start["A2"] = (
            f"Normalized-concept visual layer with {len(characters)} citation-backed character evidence records, {len(relationships)} "
            f"character-only evidence relationships, and {len(audits)} scoped source passes; "
            f"focused independent review is recorded for {review_index.get('reviewed_source_count', 0)} sources and remains pending for "
            f"{review_index.get('pending_source_count', len(audits))}."
        )
        start["A2"].font = Font(color=TEAL, italic=True, size=10)
        start["A3"] = "A normalized-concept atlas: being/entity and role/vocation concepts are visual marks; characters and source-native terms are cited evidence"
        start["B5"] = "4.0"
        start["A6"] = "Source fingerprint"
        start["B6"] = fingerprint

    workbook.properties.title = "Fantasy & High-Fantasy Comparative Archetype Atlas — Version 4.0"
    workbook.properties.subject = "Normalized-concept fantasy atlas with citation-backed character and source-term evidence"
    workbook.properties.description = (
        "Version 4 renders normalized being/entity and role/vocation concepts as graph nodes. "
        "Characters and source-native terms remain citation-backed evidence records."
    )
    workbook.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True, calcOnSave=True)
    workbook.save(args.output)

    office_binary = shutil.which("libreoffice") or shutil.which("soffice")
    if args.recalculate and office_binary:
        with tempfile.TemporaryDirectory(prefix="fantasy-atlas-v4-recalc-") as temporary_directory:
            subprocess.run(
                [
                    office_binary,
                    "--headless",
                    "--convert-to",
                    "xlsx",
                    "--outdir",
                    temporary_directory,
                    str(args.output),
                ],
                check=True,
                capture_output=True,
                text=True,
            )
            recalculated = Path(temporary_directory) / args.output.name
            if not recalculated.exists():
                raise SystemExit("LibreOffice completed without producing the recalculated Version 4 workbook")
            shutil.copy2(recalculated, args.output)
    elif args.recalculate:
        raise SystemExit("LibreOffice was requested but is not available")

    normalize_xlsx(args.output)
    print(
        f"Wrote {args.output.name}: {len(characters)} characters, {len(dimension_rows)} dimension values, "
        f"{len(relationships)} relationships, {len(terms)} source terms, {len(audits)} source passes."
    )


if __name__ == "__main__":
    main()
