#!/usr/bin/env python3
"""Validate and compile citation-backed character research batches.

Researchers write isolated JSON bundles under ``research/batch_*/*``. This
script enforces the character-evidence and relationship-endpoint contracts and
emits the browser payload only when every structural and evidence gate passes.
"""

from __future__ import annotations

import argparse
import copy
import json
import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

from openpyxl import load_workbook

from reproducible import source_fingerprint


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = ROOT / "fantasy_high_fantasy_archetype_atlas_v3.xlsx"
DEFAULT_RESEARCH_ROOT = ROOT / "research"
DEFAULT_OUTPUT = ROOT / "public" / "data" / "characters.json"
DEFAULT_REPORT = ROOT / "research" / "validation_report.json"
REVIEW_INDEX = ROOT / "research" / "independent_reviews" / "index.json"
DIMENSION_SCHEMA_PATH = ROOT / "research" / "dimensions.json"

REQUIRED_FILES = ("sources.json", "characters.json", "relationships.json", "source_terms.json")
DIMENSION_LABELS = json.loads(DIMENSION_SCHEMA_PATH.read_text(encoding="utf-8"))
DIMENSIONS = tuple(DIMENSION_LABELS)
RELATIONSHIP_TYPES = {
    "kin",
    "ally",
    "opponent",
    "mentor",
    "student",
    "ruler",
    "subject",
    "creator",
    "created",
    "patron",
    "champion",
    "lover",
    "spouse",
    "companion",
    "member-of-party",
    "transformed-by",
    "other",
}
CONFIDENCE_LEVELS = {"high", "moderate", "low"}
EVIDENCE_LEVELS = {
    "primary-text",
    "official-reference",
    "official-metadata",
    "publisher-metadata",
    "scholarly-reference",
}
SPOILER_LEVELS = {"none", "light", "moderate", "major"}
REVIEW_STATUSES = {"researched", "needs-review", "disputed"}
MAPPING_RELATIONS = {"exact", "close", "partial", "functional", "mechanical", "visual", "none"}
SOURCE_TERM_RECORD_KINDS = {"source-term", "research-boundary"}
COMPLETION_STATUSES = {
    "pass-complete",
    "narrow-metadata-pass-complete",
    "evidence-insufficient-zero-character-audit",
}
CHARACTER_PASS_LABELS = {
    "pass-complete": "Pass complete",
    "narrow-metadata-pass-complete": "Limited metadata pass",
    "evidence-insufficient-zero-character-audit": "Evidence insufficient",
}
MINIMUM_SECOND_REVIEW_CLAIM_COVERAGE = 0.20
DISALLOWED_EVIDENCE_HOSTS = {
    "wikipedia.org",
    "fandom.com",
    "reddit.com",
}
WITNESS_IDENTITY_STOP_WORDS = {
    "a",
    "an",
    "and",
    "at",
    "book",
    "by",
    "chapter",
    "chapters",
    "cited",
    "edition",
    "for",
    "frame",
    "from",
    "in",
    "of",
    "official",
    "on",
    "opening",
    "page",
    "pages",
    "paragraph",
    "paragraphs",
    "part",
    "pass",
    "section",
    "sections",
    "selected",
    "source",
    "the",
    "this",
    "to",
    "vol",
    "volume",
    "witness",
    "with",
}


def workbook_records(sheet: Any, header_row: int = 1) -> list[dict[str, Any]]:
    rows = sheet.iter_rows(min_row=header_row, values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(rows)]
    output: list[dict[str, Any]] = []
    for row in rows:
        record = {
            header: value
            for header, value in zip(headers, row)
            if header and value not in (None, "")
        }
        if record:
            output.append(record)
    return output


def load_array(path: Path, errors: list[str]) -> list[dict[str, Any]]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        errors.append(f"{path.relative_to(ROOT)}: cannot load JSON ({exc})")
        return []
    if not isinstance(value, list):
        errors.append(f"{path.relative_to(ROOT)}: root must be a JSON array")
        return []
    if not all(isinstance(item, dict) for item in value):
        errors.append(f"{path.relative_to(ROOT)}: every array item must be an object")
        return []
    return value


def require(record: dict[str, Any], keys: tuple[str, ...], context: str, errors: list[str]) -> None:
    for key in keys:
        # Presence and type are validated separately where an empty array or
        # intentionally blank caution/note is a meaningful, honest value.
        if key not in record or record[key] is None:
            errors.append(f"{context}: missing required field {key}")


def host_is_disallowed(url: str) -> bool:
    host = (urlparse(url).hostname or "").lower()
    return any(host == blocked or host.endswith(f".{blocked}") for blocked in DISALLOWED_EVIDENCE_HOSTS)


def validate_citations(
    citations: Any,
    context: str,
    errors: list[str],
    warnings: list[str],
) -> None:
    if not isinstance(citations, list) or not citations:
        errors.append(f"{context}: at least one claim-specific citation is required")
        return
    allowed_count = 0
    for index, citation in enumerate(citations, 1):
        item_context = f"{context} citation {index}"
        if not isinstance(citation, dict):
            errors.append(f"{item_context}: citation must be an object")
            continue
        url = str(citation.get("url", "")).strip()
        locator = str(citation.get("locator", "")).strip()
        supports = citation.get("supports")
        if not re.match(r"^https?://", url):
            errors.append(f"{item_context}: stable http(s) URL required")
        elif host_is_disallowed(url):
            warnings.append(f"{item_context}: orientation-only host cannot establish canonical evidence ({url})")
        else:
            allowed_count += 1
        if not locator:
            errors.append(f"{item_context}: work/chapter/page/section locator required")
        if not isinstance(supports, list) or not supports or not all(str(value).strip() for value in supports):
            errors.append(f"{item_context}: supports must be a non-empty list of claims")
    if allowed_count == 0:
        errors.append(f"{context}: no allowable canonical evidence source remains")


def source_sort_key(source_id: str) -> tuple[int, str]:
    match = re.search(r"(\d+)$", source_id)
    return (int(match.group(1)) if match else 999999, source_id)


def fold_identity(value: Any) -> str:
    normalized = unicodedata.normalize("NFKD", str(value or ""))
    return "".join(character for character in normalized.casefold() if character.isalnum())


def identity_tokens(value: Any) -> tuple[str, ...]:
    normalized = unicodedata.normalize("NFKD", str(value or "")).casefold()
    without_marks = "".join(character for character in normalized if not unicodedata.combining(character))
    return tuple(fold_identity(token) for token in re.findall(r"[^\W_]+", without_marks) if fold_identity(token))


def contains_identity_tokens(value: Any, candidate: Any) -> bool:
    value_tokens = identity_tokens(value)
    candidate_tokens = identity_tokens(candidate)
    return bool(candidate_tokens) and any(
        value_tokens[index : index + len(candidate_tokens)] == candidate_tokens
        for index in range(len(value_tokens) - len(candidate_tokens) + 1)
    )


def source_work_names(audit: dict[str, Any]) -> list[str]:
    names: list[str] = []
    for witness in audit.get("work_or_witnesses", []):
        if isinstance(witness, str):
            name = witness.strip()
        elif isinstance(witness, dict):
            name = str(
                witness.get("work")
                or witness.get("work_or_witness")
                or witness.get("witness")
                or ""
            ).strip()
        else:
            name = ""
        if name:
            names.append(name)
    return names


def source_work_entries(audit: dict[str, Any]) -> list[tuple[str, str]]:
    entries: list[tuple[str, str]] = []
    for witness in audit.get("work_or_witnesses", []):
        if isinstance(witness, str):
            name = witness.strip()
            url = ""
        elif isinstance(witness, dict):
            name = str(
                witness.get("work")
                or witness.get("work_or_witness")
                or witness.get("witness")
                or ""
            ).strip()
            url = str(witness.get("url", "")).strip()
        else:
            name = ""
            url = ""
        if name:
            entries.append((name, url))
    return entries


def urls_identify_same_witness(left: Any, right: Any) -> bool:
    left_url = urlparse(str(left or "").strip())
    right_url = urlparse(str(right or "").strip())
    if not left_url.netloc or left_url.netloc.casefold() != right_url.netloc.casefold():
        return False
    left_location = (left_url.path.rstrip("/"), left_url.query)
    right_location = (right_url.path.rstrip("/"), right_url.query)
    if left_location == right_location:
        return True
    left_document_ids = {
        token.casefold()
        for token in re.findall(r"[A-Za-z]+\d{4,}[A-Za-z0-9]*", left_url.path)
    }
    right_document_ids = {
        token.casefold()
        for token in re.findall(r"[A-Za-z]+\d{4,}[A-Za-z0-9]*", right_url.path)
    }
    return bool(left_document_ids & right_document_ids)


def urls_share_host(left: Any, right: Any) -> bool:
    left_host = (urlparse(str(left or "").strip()).hostname or "").casefold()
    right_host = (urlparse(str(right or "").strip()).hostname or "").casefold()
    return bool(left_host) and left_host == right_host


def witness_identity_score(candidate: Any, evidence: Any) -> int:
    candidate_tokens = {
        token
        for token in identity_tokens(candidate)
        if token not in WITNESS_IDENTITY_STOP_WORDS and (len(token) >= 3 or any(character.isdigit() for character in token))
    }
    evidence_tokens = {
        token
        for token in identity_tokens(evidence)
        if token not in WITNESS_IDENTITY_STOP_WORDS
    }
    return len(candidate_tokens & evidence_tokens)


def is_specific_claimed_identity(value: Any, audit: dict[str, Any]) -> bool:
    if fold_identity(value) in {
        fold_identity(audit.get("source_title", "")),
        fold_identity(f"About {audit.get('source_title', '')}"),
    }:
        return False
    meaningful_tokens = [
        token
        for token in identity_tokens(value)
        if token not in WITNESS_IDENTITY_STOP_WORDS
        and len(token) >= 3
        and not token.isdigit()
    ]
    return len(meaningful_tokens) >= 2


def audit_reference_locators(citation: dict[str, Any], audit: dict[str, Any]) -> list[str]:
    citation_url = citation.get("url", "")
    references = [
        reference
        for collection in (audit.get("citations", []), audit.get("sources", []))
        for reference in collection or []
        if isinstance(reference, dict)
    ]
    exact = [
        str(reference.get("locator", "")).strip()
        for reference in references
        if urls_identify_same_witness(citation_url, reference.get("url", ""))
        and str(reference.get("locator", "")).strip()
    ]
    if exact:
        return exact
    return [
        str(reference.get("locator", "")).strip()
        for reference in references
        if urls_share_host(citation_url, reference.get("url", ""))
        and str(reference.get("locator", "")).strip()
    ]


def claimed_witness_segment(
    record: dict[str, Any],
    audit: dict[str, Any],
    citation_index: int,
) -> str:
    explicit = str(record.get("witness_identity", "")).strip()
    if explicit:
        return explicit
    work_or_witness = str(record.get("work_or_witness", "")).strip()
    citations = record.get("citations", [])
    segments = [segment.strip() for segment in work_or_witness.split(";") if segment.strip()]
    segment = segments[citation_index] if len(segments) == len(citations) else work_or_witness
    identity, separator, locator = segment.partition(" · ")
    source_title = str(audit.get("source_title", "")).strip()
    if separator and source_title and fold_identity(identity) == fold_identity(source_title):
        return locator.strip()
    return identity.strip()


def source_work_descriptions(audit: dict[str, Any]) -> list[tuple[str, str]]:
    descriptions: list[tuple[str, str]] = []
    for witness in audit.get("work_or_witnesses", []):
        if isinstance(witness, str):
            name = witness.strip()
            description = name
        elif isinstance(witness, dict):
            name = str(
                witness.get("work")
                or witness.get("work_or_witness")
                or witness.get("witness")
                or ""
            ).strip()
            description = " ".join(
                value
                for value in (name, str(witness.get("edition", "")).strip())
                if value
            )
        else:
            name = ""
            description = ""
        if name:
            descriptions.append((name, description))
    return descriptions


def audit_witness_for_citation(
    citation: dict[str, Any],
    audit: dict[str, Any],
    claimed_identity: str,
) -> str:
    entries = source_work_entries(audit)
    citation_url = citation.get("url", "")
    url_matches = {
        name
        for name, witness_url in entries
        if witness_url and urls_identify_same_witness(citation_url, witness_url)
    }
    if len(url_matches) == 1:
        matched_name = next(iter(url_matches))
        if is_specific_claimed_identity(claimed_identity, audit) and witness_identity_score(
            matched_name, claimed_identity
        ):
            return claimed_identity
        return matched_name
    host_matches = {
        name
        for name, witness_url in entries
        if witness_url and urls_share_host(citation_url, witness_url)
    }
    if len(host_matches) == 1:
        matched_name = next(iter(host_matches))
        if is_specific_claimed_identity(claimed_identity, audit) and witness_identity_score(
            matched_name, claimed_identity
        ):
            return claimed_identity
        return matched_name
    audit_citations = audit.get("citations", [])
    citation_indexes = [
        index
        for index, audit_citation in enumerate(audit_citations)
        if isinstance(audit_citation, dict)
        and urls_identify_same_witness(citation_url, audit_citation.get("url", ""))
    ]
    if len(entries) == len(audit_citations) and len(citation_indexes) == 1:
        return entries[citation_indexes[0]][0]
    reference_locators = audit_reference_locators(citation, audit)
    if reference_locators:
        evidence = " ".join([claimed_identity, citation_url, *reference_locators])
        scored_names = [
            (witness_identity_score(name, evidence), name)
            for name in source_work_names(audit)
        ]
        best_score = max((score for score, _ in scored_names), default=0)
        best_names = [name for score, name in scored_names if score == best_score and score > 0]
        if len(best_names) == 1:
            if is_specific_claimed_identity(claimed_identity, audit) and witness_identity_score(
                best_names[0], claimed_identity
            ):
                return claimed_identity
            return best_names[0]
        source_title = str(audit.get("source_title", "")).strip()
        if claimed_identity and fold_identity(claimed_identity) not in {
            fold_identity(source_title),
            fold_identity(f"About {source_title}"),
        }:
            return claimed_identity
        if len(reference_locators) == 1:
            return reference_locators[0]
    declared_matches = [
        name
        for name in source_work_names(audit)
        if contains_identity_tokens(claimed_identity, name)
        or contains_identity_tokens(name, claimed_identity)
    ]
    if declared_matches:
        return min(declared_matches, key=lambda name: (len(identity_tokens(name)), name))
    scored_descriptions = [
        (witness_identity_score(description, claimed_identity), name)
        for name, description in source_work_descriptions(audit)
    ]
    best_score = max((score for score, _ in scored_descriptions), default=0)
    best_names = [name for score, name in scored_descriptions if score == best_score and score > 0]
    if len(best_names) == 1:
        source_title = str(audit.get("source_title", "")).strip()
        if claimed_identity and fold_identity(claimed_identity) != fold_identity(source_title):
            return claimed_identity
        return best_names[0]
    if len(entries) == 1:
        return entries[0][0]
    return ""


def source_work_name_sequences(audit: dict[str, Any]) -> list[str]:
    names = source_work_names(audit)
    sequences = ["; ".join(names)] if len(names) > 1 else []
    sequences.extend(
        name
        for name in names
        if re.search(r"\b(?:cited|read|consulted) in this pass\b", name, re.IGNORECASE)
    )
    return sequences


def claim_witness_identity(record: dict[str, Any], audit: dict[str, Any]) -> str:
    citations = record.get("citations", [])
    resolved = [
        audit_witness_for_citation(citation, audit, claimed_witness_segment(record, audit, index))
        for index, citation in enumerate(citations)
        if isinstance(citation, dict)
    ]
    return "; ".join(dict.fromkeys(resolved)) if resolved and all(resolved) else ""


def is_research_boundary_record(record: dict[str, Any]) -> bool:
    return record.get("record_kind") == "research-boundary"


def claim_ids_for_character(record: dict[str, Any]) -> list[str]:
    character_id = str(record.get("character_id", ""))
    return [
        f"character:{character_id}",
        *[
            f"dimension:{character_id}:{dimension}"
            for dimension, values in record.get("dimensions", {}).items()
            if values
        ],
    ]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--research-root", type=Path, default=DEFAULT_RESEARCH_ROOT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    args = parser.parse_args()

    workbook = load_workbook(args.workbook, read_only=True, data_only=True)
    source_rows = workbook_records(workbook["Source Corpus"])
    taxonomy_rows = workbook_records(workbook["Master Taxonomy"])
    sources_by_id = {str(row["Source_ID"]): row for row in source_rows}
    archetypes_by_id = {str(row["Archetype_ID"]): row for row in taxonomy_rows}

    discovered_bundle_dirs = sorted(
        path
        for path in args.research_root.glob("batch_*/*")
        if path.is_dir() and any((path / filename).exists() for filename in REQUIRED_FILES)
    )
    quarantined_bundle_dirs = [path for path in discovered_bundle_dirs if (path / "QUARANTINED.md").exists()]
    bundle_dirs = [path for path in discovered_bundle_dirs if path not in quarantined_bundle_dirs]
    errors: list[str] = []
    warnings: list[str] = []
    all_sources: list[dict[str, Any]] = []
    all_characters: list[dict[str, Any]] = []
    all_relationships: list[dict[str, Any]] = []
    all_term_records: list[dict[str, Any]] = []

    if not bundle_dirs:
        errors.append("No research bundles found under research/batch_*/*")

    for bundle in bundle_dirs:
        missing = [filename for filename in REQUIRED_FILES if not (bundle / filename).exists()]
        if missing:
            errors.append(f"{bundle.relative_to(ROOT)}: missing {', '.join(missing)}")
            continue
        all_sources.extend(load_array(bundle / "sources.json", errors))
        all_characters.extend(load_array(bundle / "characters.json", errors))
        all_relationships.extend(load_array(bundle / "relationships.json", errors))
        all_term_records.extend(load_array(bundle / "source_terms.json", errors))

    for record in all_term_records:
        if record.get("record_kind") not in SOURCE_TERM_RECORD_KINDS:
            errors.append(
                f"source term {record.get('term_id', '<unknown>')}: invalid record_kind {record.get('record_kind')}"
            )
    all_boundaries = [record for record in all_term_records if is_research_boundary_record(record)]
    all_terms = [record for record in all_term_records if not is_research_boundary_record(record)]
    duplicate_term_record_ids = [
        value
        for value, count in Counter(str(record.get("term_id", "")) for record in all_term_records).items()
        if value and count > 1
    ]
    if duplicate_term_record_ids:
        errors.append(f"Duplicate source term record IDs: {', '.join(sorted(duplicate_term_record_ids))}")

    source_audit_ids: list[str] = []
    for record in all_sources:
        source_id = str(record.get("source_id", ""))
        context = f"source audit {source_id or '<unknown>'}"
        require(
            record,
            (
                "source_id",
                "source_title",
                "continuity_scope",
                "work_or_witnesses",
                "coverage_rule",
                "completion_status",
                "evidence_basis",
                "citations",
                "last_reviewed",
            ),
            context,
            errors,
        )
        if source_id not in sources_by_id:
            errors.append(f"{context}: unknown Source_ID")
        source_audit_ids.append(source_id)
        if not isinstance(record.get("work_or_witnesses"), list):
            errors.append(f"{context}: work_or_witnesses must be an array")
        for field in ("in_scope_character_count", "completed_character_count"):
            if not isinstance(record.get(field), int) or int(record.get(field, -1)) < 0:
                errors.append(f"{context}: {field} must be a non-negative integer")
        if record.get("completion_status") == "pass-complete" and record.get("in_scope_character_count") != record.get("completed_character_count"):
            errors.append(f"{context}: pass-complete count does not match declared in-scope count")
        for field in ("omissions", "uncertainties"):
            if not isinstance(record.get(field, []), list):
                errors.append(f"{context}: {field} must be an array")
        validate_citations(record.get("citations"), context, errors, warnings)

    duplicate_source_audits = [value for value, count in Counter(source_audit_ids).items() if value and count > 1]
    if duplicate_source_audits:
        errors.append(f"Duplicate source audits: {', '.join(sorted(duplicate_source_audits))}")
    missing_source_audits = sorted(set(sources_by_id) - set(source_audit_ids), key=source_sort_key)
    if missing_source_audits:
        errors.append(f"Missing source audits: {', '.join(missing_source_audits)}")
    audits_by_source = {str(record.get("source_id")): record for record in all_sources}
    for record in all_sources:
        status = str(record.get("completion_status", "")).strip().lower()
        if status not in COMPLETION_STATUSES:
            errors.append(
                f"source audit {record.get('source_id', '<unknown>')}: "
                f"invalid completion_status {record.get('completion_status')}"
            )

    character_ids: list[str] = []
    character_source_by_id: dict[str, str] = {}
    character_count_by_source: Counter[str] = Counter()
    evidence_urls_by_source: dict[str, set[str]] = defaultdict(set)
    evidence_locators_by_source: dict[str, set[str]] = defaultdict(set)
    for record in all_characters:
        character_id = str(record.get("character_id", ""))
        source_id = str(record.get("source_id", ""))
        context = f"character {character_id or '<unknown>'}"
        require(
            record,
            (
                "character_id",
                "canonical_name",
                "source_id",
                "source_title",
                "continuity",
                "work_or_witness",
                "character_kind",
                "description",
                "dimensions",
                "canon_status",
                "evidence_level",
                "citations",
                "comparison_cautions",
                "spoiler_level",
                "review_status",
            ),
            context,
            errors,
        )
        if not re.match(r"^CHR-SRC\d{3}-\d{3,}$", character_id):
            errors.append(f"{context}: ID must match CHR-SRC###-###")
        if source_id not in sources_by_id:
            errors.append(f"{context}: unknown Source_ID {source_id}")
        expected_prefix = f"CHR-{source_id.replace('SRC-', 'SRC')}-" if source_id else ""
        if expected_prefix and not character_id.startswith(expected_prefix):
            errors.append(f"{context}: ID prefix does not match source {source_id}")
        if record.get("character_kind") != "individual":
            errors.append(
                f"{context}: character_kind must be individual; "
                "collectives, manifestations, generic units, and other concepts belong in source_terms"
            )
        character_ids.append(character_id)
        character_source_by_id[character_id] = source_id
        character_count_by_source[source_id] += 1
        if not isinstance(record.get("aliases", []), list):
            errors.append(f"{context}: aliases must be an array")
        if not isinstance(record.get("comparison_cautions"), list):
            errors.append(f"{context}: comparison_cautions must be an array")
        if record.get("evidence_level") not in EVIDENCE_LEVELS:
            errors.append(f"{context}: invalid evidence_level {record.get('evidence_level')}")
        if record.get("spoiler_level") not in SPOILER_LEVELS:
            errors.append(f"{context}: invalid spoiler_level {record.get('spoiler_level')}")
        if record.get("review_status") not in REVIEW_STATUSES:
            errors.append(f"{context}: invalid review_status {record.get('review_status')}")
        dimensions = record.get("dimensions")
        if not isinstance(dimensions, dict):
            errors.append(f"{context}: dimensions must be an object")
        else:
            unknown_dimensions = sorted(set(dimensions) - set(DIMENSIONS))
            if unknown_dimensions:
                errors.append(f"{context}: unknown dimensions {', '.join(unknown_dimensions)}")
            for dimension in DIMENSIONS:
                values = dimensions.get(dimension, [])
                if not isinstance(values, list):
                    errors.append(f"{context}: dimension {dimension} must be an array")
                    continue
                for index, value in enumerate(values, 1):
                    value_context = f"{context} {dimension}[{index}]"
                    if not isinstance(value, dict):
                        errors.append(f"{value_context}: dimension value must be an object")
                        continue
                    require(value, ("term", "archetype_ids", "confidence", "note"), value_context, errors)
                    if value.get("confidence") not in CONFIDENCE_LEVELS:
                        errors.append(f"{value_context}: invalid confidence {value.get('confidence')}")
                    archetype_ids = value.get("archetype_ids", [])
                    if not isinstance(archetype_ids, list):
                        errors.append(f"{value_context}: archetype_ids must be an array")
                    else:
                        for archetype_id in archetype_ids:
                            if archetype_id not in archetypes_by_id:
                                errors.append(f"{value_context}: unknown archetype ID {archetype_id}")
        validate_citations(record.get("citations"), context, errors, warnings)
        for citation in record.get("citations", []):
            if isinstance(citation, dict):
                evidence_urls_by_source[source_id].add(str(citation.get("url", "")))
                evidence_locators_by_source[source_id].add(str(citation.get("locator", "")))

    duplicate_characters = [value for value, count in Counter(character_ids).items() if value and count > 1]
    if duplicate_characters:
        errors.append(f"Duplicate character IDs: {', '.join(sorted(duplicate_characters))}")
    character_id_set = set(character_ids)

    for source_id, count in character_count_by_source.items():
        unique_urls = len(evidence_urls_by_source[source_id])
        unique_locators = len(evidence_locators_by_source[source_id])
        if count >= 10 and unique_urls <= 1:
            warnings.append(
                f"{source_id}: {count} character records depend on one evidence URL; prioritize independent source-to-claim review"
            )
        if count >= 10 and unique_locators / count < 0.3:
            warnings.append(
                f"{source_id}: only {unique_locators} unique locators for {count} characters; check claim granularity"
            )

    relationship_ids: list[str] = []
    for record in all_relationships:
        relationship_id = str(record.get("relationship_id", ""))
        context = f"relationship {relationship_id or '<unknown>'}"
        require(
            record,
            (
                "relationship_id",
                "source_id",
                "source_character_id",
                "target_character_id",
                "relationship_type",
                "label",
                "direction",
                "continuity",
                "citations",
                "confidence",
                "note",
            ),
            context,
            errors,
        )
        relationship_ids.append(relationship_id)
        if not re.match(r"^REL-SRC\d{3}-\d{3,}$", relationship_id):
            errors.append(f"{context}: ID must match REL-SRC###-###")
        source_id = str(record.get("source_id", ""))
        expected_prefix = f"REL-{source_id.replace('SRC-', 'SRC')}-" if source_id else ""
        if expected_prefix and not relationship_id.startswith(expected_prefix):
            errors.append(f"{context}: ID prefix does not match source {source_id}")
        for endpoint in ("source_character_id", "target_character_id"):
            if record.get(endpoint) not in character_id_set:
                errors.append(f"{context}: {endpoint} is not an included character ({record.get(endpoint)})")
            elif character_source_by_id.get(str(record.get(endpoint))) != source_id:
                errors.append(f"{context}: {endpoint} belongs to a different source pass")
        if record.get("source_character_id") == record.get("target_character_id"):
            errors.append(f"{context}: self-loop relationships are not allowed")
        if record.get("relationship_type") not in RELATIONSHIP_TYPES:
            errors.append(f"{context}: invalid relationship_type {record.get('relationship_type')}")
        if record.get("direction") not in {"directed", "undirected"}:
            errors.append(f"{context}: direction must be directed or undirected")
        if record.get("confidence") not in CONFIDENCE_LEVELS:
            errors.append(f"{context}: invalid confidence {record.get('confidence')}")
        if record.get("source_id") not in sources_by_id:
            errors.append(f"{context}: unknown Source_ID {record.get('source_id')}")
        validate_citations(record.get("citations"), context, errors, warnings)

    duplicate_relationships = [value for value, count in Counter(relationship_ids).items() if value and count > 1]
    if duplicate_relationships:
        errors.append(f"Duplicate relationship IDs: {', '.join(sorted(duplicate_relationships))}")

    character_terms_by_source_dimension: dict[tuple[str, str], dict[str, str]] = defaultdict(dict)
    for character in all_characters:
        for dimension, values in character.get("dimensions", {}).items():
            for value in values:
                term = str(value.get("term", "")).strip()
                if term:
                    character_terms_by_source_dimension[(str(character.get("source_id", "")), dimension)][
                        fold_identity(term)
                    ] = term

    term_count_by_source: Counter[str] = Counter()
    term_witness_identities: dict[str, str] = {}
    for record in all_terms:
        term_id = str(record.get("term_id", ""))
        context = f"source term {term_id or '<unknown>'}"
        require(
            record,
            (
                "term_id",
                "record_kind",
                "source_id",
                "canonical_term",
                "work_or_witness",
                "dimension",
                "archetype_ids",
                "mapping_relation",
                "definition",
                "cultural_caution",
                "citations",
                "review_status",
            ),
            context,
            errors,
        )
        term_count_by_source[str(record.get("source_id", ""))] += 1
        if record.get("source_id") not in sources_by_id:
            errors.append(f"{context}: unknown Source_ID {record.get('source_id')}")
        work_or_witness = record.get("work_or_witness")
        if not isinstance(work_or_witness, str) or not work_or_witness.strip():
            errors.append(f"{context}: work_or_witness must be a non-empty string")
        elif not any(character.isalpha() for character in work_or_witness):
            errors.append(f"{context}: work_or_witness must identify the claim-specific work or witness")
        else:
            audit = audits_by_source.get(str(record.get("source_id", "")), {})
            witness_identity = claim_witness_identity(record, audit)
            if not witness_identity:
                errors.append(f"{context}: witness_identity must identify the work separately from its locator")
            else:
                term_witness_identities[term_id] = witness_identity
            if any(
                work_or_witness == sequence or work_or_witness.startswith(f"{sequence} · ")
                for sequence in source_work_name_sequences(audit)
            ):
                errors.append(f"{context}: work_or_witness must not prepend the source-wide witness scope")
            for citation_index, citation in enumerate(record.get("citations", []), 1):
                locator = str(citation.get("locator", "")).strip() if isinstance(citation, dict) else ""
                if locator and locator not in work_or_witness:
                    errors.append(
                        f"{context}: work_or_witness must pair citation {citation_index} with its claim locator"
                    )
        identity_forms = record.get("identity_forms", [])
        if not isinstance(identity_forms, list) or not all(
            isinstance(value, str) and value.strip() for value in identity_forms
        ):
            errors.append(f"{context}: identity_forms must be an array of non-empty strings when present")
            identity_forms = []
        identity_fields = [
            str(record.get(field_name, "")).strip()
            for field_name in ("canonical_term", "transliteration", "original_script")
            if str(record.get(field_name, "")).strip()
        ]
        whole_identity_forms = {fold_identity(value) for value in identity_fields}
        character_terms = character_terms_by_source_dimension.get(
            (str(record.get("source_id", "")), str(record.get("dimension", ""))),
            {},
        )
        explicit_identity_forms = {fold_identity(value) for value in identity_forms}
        missing_identity_forms = sorted(
            term
            for folded_term, term in character_terms.items()
            if folded_term not in whole_identity_forms
            and folded_term not in explicit_identity_forms
            and any(contains_identity_tokens(identity_field, term) for identity_field in identity_fields)
        )
        if missing_identity_forms:
            errors.append(
                f"{context}: character-linked compound identity forms must be explicit: "
                + ", ".join(missing_identity_forms)
            )
        if record.get("dimension") not in DIMENSIONS:
            errors.append(f"{context}: invalid dimension {record.get('dimension')}")
        archetype_ids = record.get("archetype_ids")
        if not isinstance(archetype_ids, list):
            errors.append(f"{context}: archetype_ids must be an array")
        else:
            for archetype_id in archetype_ids:
                if archetype_id not in archetypes_by_id:
                    errors.append(f"{context}: unknown archetype ID {archetype_id}")
        if record.get("mapping_relation") not in MAPPING_RELATIONS:
            errors.append(f"{context}: invalid mapping_relation {record.get('mapping_relation')}")
        if record.get("review_status") not in REVIEW_STATUSES:
            errors.append(f"{context}: invalid review_status {record.get('review_status')}")
        validate_citations(record.get("citations"), context, errors, warnings)

    for record in all_boundaries:
        boundary_id = str(record.get("term_id", ""))
        context = f"research boundary {boundary_id or '<unknown>'}"
        require(
            record,
            ("term_id", "record_kind", "source_id", "canonical_term", "definition", "citations", "review_status"),
            context,
            errors,
        )
        if record.get("source_id") not in sources_by_id:
            errors.append(f"{context}: unknown Source_ID {record.get('source_id')}")
        if record.get("review_status") not in REVIEW_STATUSES:
            errors.append(f"{context}: invalid review_status {record.get('review_status')}")
        validate_citations(record.get("citations"), context, errors, warnings)
    for source_id, audit in audits_by_source.items():
        count = character_count_by_source[source_id]
        declared = audit.get("completed_character_count")
        if declared != count:
            errors.append(f"{source_id}: audit declares {declared} completed characters but bundle contains {count}")

    dimension_term_counts: dict[str, int] = defaultdict(int)
    for record in all_characters:
        for dimension, values in record.get("dimensions", {}).items():
            dimension_term_counts[dimension] += len(values)

    review_index: dict[str, Any] = {}
    try:
        loaded_review_index = json.loads(REVIEW_INDEX.read_text(encoding="utf-8"))
        if not isinstance(loaded_review_index, dict):
            raise ValueError("root must be an object")
        review_index = loaded_review_index
    except (OSError, json.JSONDecodeError, ValueError) as exc:
        errors.append(f"{REVIEW_INDEX.relative_to(ROOT)}: cannot load independent-review index ({exc})")

    review_status_map = review_index.get("source_review_status", {})
    if review_index and not isinstance(review_status_map, dict):
        errors.append("independent-review index: source_review_status must be an object")
        review_status_map = {}
    unknown_review_sources = sorted(set(review_status_map) - set(sources_by_id), key=source_sort_key)
    if unknown_review_sources:
        errors.append(f"independent-review index: unknown source IDs {', '.join(unknown_review_sources)}")
    reviewed_source_count = len(review_status_map)
    if review_index and review_index.get("reviewed_source_count") != reviewed_source_count:
        errors.append("independent-review index: reviewed_source_count does not match source_review_status")
    if review_index and review_index.get("pending_source_count") != len(sources_by_id) - reviewed_source_count:
        errors.append("independent-review index: pending_source_count does not match corpus coverage")
    warning_source_ids = sorted(
        {warning.split(":", 1)[0] for warning in warnings if re.match(r"^SRC-\d{3}:", warning)},
        key=source_sort_key,
    )
    indexed_warning_sources = sorted(review_index.get("retained_warning_source_ids", []), key=source_sort_key)
    if review_index and indexed_warning_sources != warning_source_ids:
        errors.append("independent-review index: retained warning source IDs do not match current validator warnings")

    raw_claim_ids = {
        *[
            claim_id
            for record in all_characters
            for claim_id in claim_ids_for_character(record)
        ],
        *[f"relationship:{record.get('relationship_id', '')}" for record in all_relationships],
        *[f"source-term:{record.get('term_id', '')}" for record in all_terms],
    }
    research_boundary_review_ids = {
        f"source-term:{record.get('term_id', '')}" for record in all_boundaries
    }
    review_ledger_names = {
        str(ledger)
        for review in review_index.get("corpus_wide_reviews", [])
        if isinstance(review, dict)
        for ledger in [review.get("ledger")]
        if ledger
    }
    for review in review_status_map.values():
        if not isinstance(review, dict):
            continue
        review_ledger_names.update(str(ledger) for ledger in review.get("ledgers", []) if ledger)
    zero_character_audit = review_index.get("zero_character_audit", {})
    if isinstance(zero_character_audit, dict) and zero_character_audit.get("ledger"):
        review_ledger_names.add(str(zero_character_audit["ledger"]))
    reviewed_claim_ids: set[str] = set()
    reviewed_claim_ids_by_ledger: dict[str, set[str]] = {}
    review_ledger_paths: list[Path] = []
    for ledger_name in sorted(review_ledger_names):
        ledger_path = REVIEW_INDEX.parent / ledger_name
        review_ledger_paths.append(ledger_path)
        try:
            ledger = json.loads(ledger_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            errors.append(f"{ledger_path.relative_to(ROOT)}: cannot load review ledger ({exc})")
            continue
        entries = ledger if isinstance(ledger, list) else [ledger]
        ledger_claim_ids: set[str] = set()
        for entry in entries:
            if not isinstance(entry, dict):
                continue
            claim_ids = entry.get("reviewed_claim_ids", [])
            if not isinstance(claim_ids, list) or not all(isinstance(value, str) and value for value in claim_ids):
                errors.append(f"{ledger_path.relative_to(ROOT)}: reviewed_claim_ids must contain stable claim IDs")
                continue
            inspected_count = entry.get("inspected_count")
            if isinstance(inspected_count, int) and len(set(claim_ids)) != inspected_count:
                errors.append(
                    f"{ledger_path.relative_to(ROOT)}: reviewed_claim_ids do not match the documented inspected_count"
                )
            ledger_claim_ids.update(claim_ids)
        reviewed_claim_ids_by_ledger[ledger_name] = ledger_claim_ids
        reviewed_claim_ids.update(ledger_claim_ids)
    unknown_review_claim_ids = sorted(
        reviewed_claim_ids - raw_claim_ids - research_boundary_review_ids
    )
    if unknown_review_claim_ids:
        errors.append(
            "independent-review ledgers reference unknown claim IDs: "
            + ", ".join(unknown_review_claim_ids)
        )
    reviewed_claim_ids.intersection_update(raw_claim_ids)
    reviewed_claim_ids_by_ledger = {
        ledger_name: claim_ids & raw_claim_ids
        for ledger_name, claim_ids in reviewed_claim_ids_by_ledger.items()
    }

    unresolved_character_ids = {
        str(record.get("character_id", ""))
        for record in all_characters
        if record.get("review_status") == "disputed"
    }
    unresolved_term_ids = {
        str(record.get("term_id", ""))
        for record in all_terms
        if record.get("review_status") == "disputed"
    }
    dependent_relationship_ids = {
        str(record.get("relationship_id", ""))
        for record in all_relationships
        if record.get("source_character_id") in unresolved_character_ids
        or record.get("target_character_id") in unresolved_character_ids
    }
    required_quarantine_ids = unresolved_character_ids | unresolved_term_ids | dependent_relationship_ids
    record_quarantine = review_index.get("record_quarantine", {})
    if not isinstance(record_quarantine, dict):
        errors.append("independent-review index: record_quarantine must be an object")
        record_quarantine = {}
    configured_quarantine_ids = {
        str(value) for value in record_quarantine.get("record_ids", [])
    } if isinstance(record_quarantine.get("record_ids", []), list) else set()
    if configured_quarantine_ids != required_quarantine_ids:
        errors.append("independent-review index: record quarantine does not match unresolved disputed records")
    if required_quarantine_ids and not str(record_quarantine.get("reason", "")).strip():
        errors.append("independent-review index: record quarantine requires a containment reason")

    required_mapping_quarantine_ids = {
        *[
            f"dimension:{record.get('character_id', '')}:{dimension}"
            for record in all_characters
            if record.get("character_id") not in configured_quarantine_ids
            for dimension, values in record.get("dimensions", {}).items()
            if any(value.get("archetype_ids") for value in values)
            and (
                record.get("review_status") == "needs-review"
                or (
                    bool(record.get("comparison_cautions"))
                    and f"dimension:{record.get('character_id', '')}:{dimension}" not in reviewed_claim_ids
                )
            )
        ],
        *[
            f"source-term:{record.get('term_id', '')}"
            for record in all_terms
            if record.get("term_id") not in configured_quarantine_ids
            and bool(record.get("archetype_ids"))
            and (
                record.get("review_status") == "needs-review"
                or (
                    bool(str(record.get("cultural_caution", "")).strip())
                    and f"source-term:{record.get('term_id', '')}" not in reviewed_claim_ids
                )
            )
        ],
    }
    mapping_quarantine = review_index.get("mapping_quarantine", {})
    if not isinstance(mapping_quarantine, dict):
        errors.append("independent-review index: mapping_quarantine must be an object")
        mapping_quarantine = {}
    if mapping_quarantine.get("policy") != "unreviewed-sensitive-claim-mappings":
        errors.append("independent-review index: mapping quarantine policy is not claim-based")
    configured_mapping_quarantine_ids = required_mapping_quarantine_ids
    if required_mapping_quarantine_ids and not str(mapping_quarantine.get("reason", "")).strip():
        errors.append("independent-review index: mapping quarantine requires a containment reason")

    promoted_characters = [
        copy.deepcopy(record)
        for record in all_characters
        if record.get("character_id") not in configured_quarantine_ids
    ]
    promoted_relationships = [
        record for record in all_relationships if record.get("relationship_id") not in configured_quarantine_ids
    ]
    promoted_terms = [
        copy.deepcopy(record)
        for record in all_terms
        if record.get("term_id") not in configured_quarantine_ids
    ]
    for record in promoted_characters:
        for dimension, values in record.get("dimensions", {}).items():
            if f"dimension:{record.get('character_id', '')}:{dimension}" in configured_mapping_quarantine_ids:
                for value in values:
                    value["archetype_ids"] = []
    for record in promoted_terms:
        witness_identity = term_witness_identities.get(str(record.get("term_id", "")), "")
        record["witness_identity"] = witness_identity
        if witness_identity and not contains_identity_tokens(record.get("work_or_witness", ""), witness_identity):
            record["work_or_witness"] = f"{witness_identity} · {record.get('work_or_witness', '')}"
        if f"source-term:{record.get('term_id', '')}" in configured_mapping_quarantine_ids:
            record["archetype_ids"] = []

    promoted_character_counts = Counter(str(record.get("source_id", "")) for record in promoted_characters)
    promoted_relationship_counts = Counter(str(record.get("source_id", "")) for record in promoted_relationships)
    promoted_term_counts = Counter(str(record.get("source_id", "")) for record in promoted_terms)

    promoted_claim_ids = {
        *[
            claim_id
            for record in promoted_characters
            for claim_id in claim_ids_for_character(record)
        ],
        *[f"relationship:{record.get('relationship_id', '')}" for record in promoted_relationships],
        *[f"source-term:{record.get('term_id', '')}" for record in promoted_terms],
    }
    reviewed_promoted_claim_ids = reviewed_claim_ids & promoted_claim_ids
    reviewed_claim_count = len(reviewed_promoted_claim_ids)
    review_claim_coverage = reviewed_claim_count / len(promoted_claim_ids) if promoted_claim_ids else 0.0
    if review_claim_coverage < MINIMUM_SECOND_REVIEW_CLAIM_COVERAGE:
        errors.append(
            "independent-review index: reviewed claim sample "
            f"{review_claim_coverage:.1%} is below the required {MINIMUM_SECOND_REVIEW_CLAIM_COVERAGE:.0%}"
        )

    promoted_claim_ids_by_source: dict[str, set[str]] = defaultdict(set)
    for record in promoted_characters:
        promoted_claim_ids_by_source[str(record.get("source_id", ""))].update(claim_ids_for_character(record))
    for record in promoted_relationships:
        promoted_claim_ids_by_source[str(record.get("source_id", ""))].add(
            f"relationship:{record.get('relationship_id', '')}"
        )
    for record in promoted_terms:
        promoted_claim_ids_by_source[str(record.get("source_id", ""))].add(
            f"source-term:{record.get('term_id', '')}"
        )
    focused_reviewed_claim_ids_by_source: dict[str, set[str]] = defaultdict(set)
    for source_id, source_review in review_status_map.items():
        if not isinstance(source_review, dict):
            continue
        for ledger_name in source_review.get("ledgers", []):
            focused_reviewed_claim_ids_by_source[source_id].update(
                reviewed_claim_ids_by_ledger.get(str(ledger_name), set())
            )
        focused_reviewed_claim_ids_by_source[source_id].intersection_update(
            promoted_claim_ids_by_source[source_id]
        )
    reviewed_claim_source_count = sum(
        bool(focused_reviewed_claim_ids_by_source[source_id])
        for source_id in promoted_claim_ids_by_source
    )
    fully_reviewed_claim_source_count = sum(
        bool(claim_ids) and claim_ids <= focused_reviewed_claim_ids_by_source[source_id]
        for source_id, claim_ids in promoted_claim_ids_by_source.items()
    )

    def counted_review_lane(reviewed: int, total: int, pending_detail: str) -> dict[str, str]:
        if reviewed == 0:
            return {"status": "not-started", "label": "Not started", "detail": pending_detail}
        if reviewed < total:
            return {
                "status": "in-progress",
                "label": "In progress",
                "detail": f"{reviewed} of {total} accepted claims have recorded focused review",
            }
        return {
            "status": "reviewed",
            "label": "Reviewed",
            "detail": f"All {total} accepted claims have recorded focused review",
        }

    def evidence_review_lane(count: int, evidence_name: str, pass_name: str) -> dict[str, str]:
        if count == 0:
            return {
                "status": "not-started",
                "label": "Not started",
                "detail": f"No accepted {evidence_name} evidence or {pass_name}-pass completion declaration is recorded",
            }
        return {
            "status": "in-progress",
            "label": "Evidence recorded",
            "detail": f"{count} accepted {evidence_name} records are validated; no {pass_name}-pass completion declaration is recorded",
        }

    def review_lanes(audit: dict[str, Any]) -> dict[str, dict[str, str]]:
        source_id = str(audit.get("source_id", ""))
        completion_status = str(audit.get("completion_status", "")).strip().lower()
        completed = int(audit.get("completed_character_count", 0) or 0)
        in_scope = int(audit.get("in_scope_character_count", 0) or 0)
        source_claims = promoted_claim_ids_by_source[source_id]
        reviewed_source_claims = focused_reviewed_claim_ids_by_source[source_id]
        source_review = review_status_map.get(source_id, {})
        review_types = [
            str(value).casefold()
            for value in source_review.get("review_types", [])
        ] if isinstance(source_review, dict) else []
        continuity_review_recorded = any(
            any(keyword in review_type for keyword in ("continuity", "locator", "textual-boundary", "witness"))
            for review_type in review_types
        )
        continuity_review_claims = source_claims & {
            claim_id
            for ledger_name in source_review.get("ledgers", [])
            for claim_id in reviewed_claim_ids_by_ledger.get(str(ledger_name), set())
        } if continuity_review_recorded and isinstance(source_review, dict) else set()
        continuity_lane = counted_review_lane(
            len(continuity_review_claims),
            len(source_claims),
            "No lane-specific continuity or witness review is recorded",
        ) if continuity_review_recorded else {
            "status": "not-started",
            "label": "Not started",
            "detail": "No lane-specific continuity or witness review is recorded",
        }
        return {
            "scoped": {
                "status": "pass-complete",
                "label": "Pass complete",
                "detail": str(audit.get("continuity_scope", "")),
            },
            "characterPass": {
                "status": completion_status,
                "label": CHARACTER_PASS_LABELS.get(completion_status, "Status not recognized"),
                "detail": f"{promoted_character_counts[source_id]} accepted records from {completed} of {in_scope} completed in scope",
            },
            "terminologyPass": evidence_review_lane(
                promoted_term_counts[source_id], "source-term", "terminology"
            ),
            "relationshipPass": evidence_review_lane(
                promoted_relationship_counts[source_id], "relationship", "relationship"
            ),
            "secondReview": counted_review_lane(
                len(reviewed_source_claims),
                len(source_claims),
                "No accepted claim IDs have recorded focused second review",
            ),
            "continuityReview": continuity_lane,
        }

    input_paths = [
        args.workbook,
        DIMENSION_SCHEMA_PATH,
        REVIEW_INDEX,
        *review_ledger_paths,
        Path(__file__),
        *discovered_bundle_dirs,
    ]
    fingerprint = source_fingerprint([path for path in input_paths if path.exists()])

    report = {
        "source_fingerprint": fingerprint,
        "status": "valid" if not errors else "invalid",
        "bundles": [str(path.relative_to(ROOT)) for path in bundle_dirs],
        "quarantinedBundles": [str(path.relative_to(ROOT)) for path in quarantined_bundle_dirs],
        "counts": {
            "source_audits": len(all_sources),
            "characters": len(all_characters),
            "relationships": len(all_relationships),
            "source_terms": len(all_terms),
            "research_boundaries": len(all_boundaries),
            "errors": len(errors),
            "warnings": len(warnings),
        },
        "character_counts_by_source": dict(sorted(character_count_by_source.items(), key=lambda item: source_sort_key(item[0]))),
        "dimension_value_counts": dict(sorted(dimension_term_counts.items())),
        "evidence_concentration_by_source": {
            source_id: {
                "character_records": character_count_by_source[source_id],
                "unique_character_citation_urls": len(evidence_urls_by_source[source_id]),
                "unique_character_citation_locators": len(evidence_locators_by_source[source_id]),
            }
            for source_id in sorted(character_count_by_source, key=source_sort_key)
        },
        "independent_review": {
            "reviewed_sources": fully_reviewed_claim_source_count,
            "pending_sources": len(sources_by_id) - fully_reviewed_claim_source_count,
            "retained_warning_sources": len(warning_source_ids),
            "corpus_wide_reviews": len(review_index.get("corpus_wide_reviews", [])),
            "reviewed_claims": reviewed_claim_count,
            "total_claims": len(promoted_claim_ids),
            "claim_coverage": round(review_claim_coverage, 6),
            "quarantined_records": len(configured_quarantine_ids),
            "quarantined_mappings": len(configured_mapping_quarantine_ids),
        },
        "errors": errors,
        "warnings": warnings,
    }
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    if errors:
        print(f"Character research validation failed with {len(errors)} error(s); see {args.report}")
        raise SystemExit(1)

    compiled = {
        "meta": {
            "title": "Fantasy Character Atlas",
            "version": "4.0-research",
            "sourceFingerprint": fingerprint,
            "sourceWorkbook": args.workbook.name,
            "counts": {
                **report["counts"],
                "characters": len(promoted_characters),
                "relationships": len(promoted_relationships),
                "source_terms": len(promoted_terms),
                "quarantined_records": len(configured_quarantine_ids),
                "quarantined_mappings": len(configured_mapping_quarantine_ids),
            },
            "researchCoverage": {
                "corpusSources": len(source_rows),
                "characterResearchedSources": len(all_sources),
                "characterUnresearchedSources": len(source_rows) - len(all_sources),
            },
            "qualityWarnings": warnings,
            "reviewCoverage": {
                "focusedReviewedSources": reviewed_claim_source_count,
                "pendingFullSecondReviewSources": len(sources_by_id) - fully_reviewed_claim_source_count,
                "retainedWarningSources": len(warning_source_ids),
                "retainedWarningFlags": len(warnings),
                "corpusWideContractReviews": len(review_index.get("corpus_wide_reviews", [])),
                "reviewedClaims": reviewed_claim_count,
                "totalClaims": len(promoted_claim_ids),
                "claimCoverage": round(review_claim_coverage, 6),
                "minimumClaimCoverage": MINIMUM_SECOND_REVIEW_CLAIM_COVERAGE,
                "quarantinedRecordIds": sorted(configured_quarantine_ids),
                "quarantinedMappingClaimIds": sorted(configured_mapping_quarantine_ids),
            },
            "quarantinedBundles": [str(path.relative_to(ROOT)) for path in quarantined_bundle_dirs],
            "visualContract": "Characters and source-native terms are searchable evidence; normalized beings and classes remain the graph nodes.",
        },
        "dimensions": list(DIMENSIONS),
        "dimensionLabels": DIMENSION_LABELS,
        "corpusSources": [
            {
                "sourceId": row.get("Source_ID", ""),
                "title": row.get("Title_or_Franchise", ""),
                "creator": row.get("Creator_Studio", ""),
                "medium": row.get("Medium", ""),
                "region": row.get("Region_Tradition", ""),
                "firstAppearance": row.get("First_Appearance", ""),
                "continuityUnit": row.get("Continuity_Unit", ""),
                "priorityTier": row.get("Priority_Tier", ""),
                "distinctiveContribution": row.get("Distinctive_Contribution", ""),
                "scopeNote": row.get("Scope_Note", ""),
                "referenceUrl": row.get("Reference_URL", ""),
                "status": row.get("Status", ""),
            }
            for row in sorted(source_rows, key=lambda row: source_sort_key(str(row.get("Source_ID", ""))))
        ],
        "sources": sorted(
            [
                {
                    **row,
                    "review_lanes": review_lanes(row),
                    "independent_review": review_status_map.get(
                        str(row.get("source_id", "")),
                        {
                            "review_types": [],
                            "ledgers": [],
                            "outcomes": [],
                            "unresolved_limits": ["Focused independent second review remains pending."],
                        },
                    ),
                }
                for row in all_sources
            ],
            key=lambda row: source_sort_key(str(row.get("source_id", ""))),
        ),
        "characters": sorted(promoted_characters, key=lambda row: (source_sort_key(str(row.get("source_id", ""))), str(row.get("canonical_name", "")))),
        "relationships": promoted_relationships,
        "sourceTerms": promoted_terms,
        "researchBoundaries": sorted(
            all_boundaries,
            key=lambda row: (
                source_sort_key(str(row.get("source_id", ""))),
                str(row.get("term_id", "")),
            ),
        ),
        "independentReviews": review_index,
        "taxonomy": {
            archetype_id: {
                "name": row.get("Preferred_Name", archetype_id),
                "domain": row.get("Domain", ""),
                "parentId": row.get("Parent_ID", ""),
                "tier": row.get("Tier", ""),
            }
            for archetype_id, row in archetypes_by_id.items()
        },
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(compiled, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(
        "Validated "
        f"{len(promoted_characters)} characters, {len(promoted_relationships)} character relationships, "
        f"and {len(promoted_terms)} source terms across {len(all_sources)} source passes "
        f"({len(configured_quarantine_ids)} unresolved records quarantined)."
    )


if __name__ == "__main__":
    main()
