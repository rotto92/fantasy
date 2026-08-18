#!/usr/bin/env python3
"""Extract the atlas workbook into a validated, visualization-ready graph.

The workbook remains the editorial source of truth. This script creates a
deterministic JSON projection for the browser explorer; it does not mutate the
workbook.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from reproducible import source_fingerprint, target_for_treatment


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = ROOT / "fantasy_high_fantasy_archetype_atlas_v3.xlsx"
DEFAULT_OUTPUT = ROOT / "generated" / "atlas.json"

DOMAIN_COLORS = {
    "People & Beings": "#66d9c6",
    "Roles & Vocations": "#f2b35d",
    "Power Traditions": "#a78bfa",
    "States & Transformations": "#fb7185",
    "Institutions & Affiliations": "#60a5fa",
    "Narrative Archetypes": "#f472b6",
    "Civilization & Culture Archetypes": "#facc15",
    "Game-Mechanical Archetypes": "#38bdf8",
    "Artifacts, Relics, Weapons & Vehicles": "#f97316",
    "Cosmologies, Realms & Sacred Geography": "#818cf8",
    "Metaphysical Laws, Rituals & Exchanges": "#34d399",
}

COVERAGE_COLORS = {
    "uncharted": "#536078",
    "scouted": "#a78bfa",
    "mapped": "#f2b35d",
    "target-reached": "#66d9c6",
}


def clean(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str):
        return re.sub(r"\s+", " ", value).strip()
    return value


def records(sheet: Any, header_row: int) -> list[dict[str, Any]]:
    rows = sheet.iter_rows(min_row=header_row, values_only=True)
    headers = [clean(value) for value in next(rows)]
    output: list[dict[str, Any]] = []
    for row in rows:
        item = {
            str(header): clean(value)
            for header, value in zip(headers, row)
            if header and clean(value) != ""
        }
        if item:
            output.append(item)
    return output


def split_ids(value: Any) -> list[str]:
    if not value:
        return []
    return [part.strip() for part in re.split(r"[;,|]", str(value)) if part.strip()]


def hash_fraction(text: str, salt: str = "") -> float:
    digest = hashlib.sha256(f"{salt}:{text}".encode("utf-8")).digest()
    return int.from_bytes(digest[:8], "big") / (2**64 - 1)


def coverage_level(count: int, target: int) -> str:
    if count == 0:
        return "uncharted"
    ratio = count / max(target, 1)
    if ratio < 0.4:
        return "scouted"
    if ratio < 1:
        return "mapped"
    return "target-reached"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    values = load_workbook(args.input, read_only=True, data_only=True)
    taxonomy = records(values["Master Taxonomy"], 1)
    sources = records(values["Source Corpus"], 1)
    entries = records(values["Seed Catalogue"], 1)
    priorities = records(values["Source Priority"], 6)
    terms = records(values["Term Metadata"], 4)
    crosswalks = records(values["Adaptation Crosswalk"], 4)
    concepts = records(values["Source Concept Map"], 1) if "Source Concept Map" in values.sheetnames else []
    source_relationships = records(values["Source Relationships"], 1) if "Source Relationships" in values.sheetnames else []

    archetype_by_id = {row["Archetype_ID"]: row for row in taxonomy}
    source_by_id = {row["Source_ID"]: row for row in sources}
    priority_by_source = {row["Source_ID"]: row for row in priorities if row.get("Source_ID")}
    terms_by_entry: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for term in terms:
        terms_by_entry[term.get("Source_Entry_ID", "")].append(term)

    errors: list[str] = []
    if len(archetype_by_id) != len(taxonomy):
        errors.append("Duplicate Archetype_ID values detected")
    if len(source_by_id) != len(sources):
        errors.append("Duplicate Source_ID values detected")

    entry_counts = Counter(row.get("Source_ID", "") for row in entries)
    source_archetype_counts: dict[tuple[str, str], int] = Counter()
    entry_archetypes: dict[str, list[str]] = {}

    for row in taxonomy:
        parent_id = row.get("Parent_ID", "")
        if parent_id and parent_id not in archetype_by_id:
            errors.append(f"{row['Archetype_ID']} has unknown parent {parent_id}")

    for row in entries:
        entry_id = row.get("Entry_ID", "")
        source_id = row.get("Source_ID", "")
        if source_id not in source_by_id:
            errors.append(f"{entry_id} has unknown source {source_id}")
        linked = [row.get("Primary_Archetype_ID", "")] + split_ids(
            row.get("Additional_Archetype_IDs", "")
        )
        linked = [archetype_id for archetype_id in linked if archetype_id]
        entry_archetypes[entry_id] = linked
        for archetype_id in linked:
            if archetype_id not in archetype_by_id:
                errors.append(f"{entry_id} has unknown archetype {archetype_id}")
            else:
                source_archetype_counts[(source_id, archetype_id)] += 1

    if errors:
        raise SystemExit("Workbook validation failed:\n- " + "\n- ".join(errors[:100]))

    domain_order = list(DOMAIN_COLORS)
    domain_centers: dict[str, tuple[float, float]] = {}
    for index, domain in enumerate(domain_order):
        angle = -math.pi / 2 + 2 * math.pi * index / len(domain_order)
        domain_centers[domain] = (51 * math.cos(angle), 51 * math.sin(angle))

    archetype_positions: dict[str, tuple[float, float]] = {}
    for row in taxonomy:
        archetype_id = row["Archetype_ID"]
        domain = row.get("Domain", "")
        center_x, center_y = domain_centers.get(domain, (0.0, 0.0))
        tier = int(row.get("Tier", 2) or 2)
        if tier == 1:
            archetype_positions[archetype_id] = (center_x, center_y)
            continue
        angle = 2 * math.pi * hash_fraction(archetype_id, "angle")
        radius_min = {2: 4.5, 3: 10.0, 4: 15.5}.get(tier, 20.0)
        radius_span = {2: 6.5, 3: 6.0, 4: 6.0}.get(tier, 7.0)
        radius = radius_min + radius_span * math.sqrt(hash_fraction(archetype_id, "radius"))
        archetype_positions[archetype_id] = (
            center_x + radius * math.cos(angle),
            center_y + radius * math.sin(angle),
        )

    source_links: dict[str, list[tuple[str, int]]] = defaultdict(list)
    for (source_id, archetype_id), count in source_archetype_counts.items():
        source_links[source_id].append((archetype_id, count))

    source_positions: dict[str, tuple[float, float]] = {}
    for source in sources:
        source_id = source["Source_ID"]
        linked = source_links.get(source_id, [])
        jitter_angle = 2 * math.pi * hash_fraction(source_id, "source-angle")
        if linked:
            domain_weights: Counter[str] = Counter()
            for archetype_id, weight in linked:
                domain_weights[str(archetype_by_id[archetype_id].get("Domain", ""))] += weight
            dominant_domain = domain_weights.most_common(1)[0][0]
            center_x, center_y = domain_centers[dominant_domain]
            # A pure centroid pulls broad franchises into one illegible knot at
            # the origin. This inner ring preserves their dominant conceptual
            # neighborhood while cross-domain links remain visible.
            x, y = center_x * 0.69, center_y * 0.69
            jitter = 4.0 + 6.0 * hash_fraction(source_id, "source-jitter")
            source_positions[source_id] = (
                x + jitter * math.cos(jitter_angle),
                y + jitter * math.sin(jitter_angle),
            )
        else:
            region = source.get("Region_Tradition", "Unspecified")
            group_angle = 2 * math.pi * hash_fraction(region, "region")
            local_angle = (hash_fraction(source_id, "frontier") - 0.5) * 0.22
            radius = 91 + 13 * hash_fraction(source_id, "frontier-radius")
            angle = group_angle + local_angle
            source_positions[source_id] = (radius * math.cos(angle), radius * math.sin(angle))

    graph_nodes: list[dict[str, Any]] = []
    graph_edges: list[dict[str, Any]] = []

    for row in taxonomy:
        archetype_id = row["Archetype_ID"]
        tier = int(row.get("Tier", 2) or 2)
        domain = row.get("Domain", "")
        x, y = archetype_positions[archetype_id]
        graph_nodes.append(
            {
                "id": archetype_id,
                "kind": "domain" if tier == 1 else "archetype",
                "label": row.get("Preferred_Name", archetype_id),
                "x": round(x, 5),
                "y": round(y, 5),
                "size": 17 if tier == 1 else max(2.0, 7.5 - tier * 1.35),
                "color": DOMAIN_COLORS.get(domain, "#94a3b8"),
                "domain": domain,
                "tier": tier,
                "record": row,
            }
        )
        parent_id = row.get("Parent_ID", "")
        if parent_id:
            graph_edges.append(
                {
                    "id": f"taxonomy:{parent_id}:{archetype_id}",
                    "source": parent_id,
                    "target": archetype_id,
                    "kind": "taxonomy",
                    "weight": 1,
                    "label": "subtype of",
                }
            )

    source_coverage: list[dict[str, Any]] = []
    for source in sources:
        source_id = source["Source_ID"]
        count = entry_counts[source_id]
        priority = priority_by_source.get(source_id, {})
        treatment = str(priority.get("Recommended_Treatment", "Representative pass"))
        target = target_for_treatment(treatment)
        level = coverage_level(count, target)
        score = priority.get("Marginal_Value_Score", priority.get("Weighted_Score", ""))
        x, y = source_positions[source_id]
        source_record = dict(source)
        source_record.update(
            {
                "Entry_Count": count,
                "Coverage_Target": target,
                "Coverage_Level": level,
                "Recommended_Treatment": treatment,
                "Marginal_Value_Score": score,
                "Baseline_Rank": priority.get("Baseline_Rank", ""),
            }
        )
        source_coverage.append(source_record)
        graph_nodes.append(
            {
                "id": source_id,
                "kind": "source",
                "label": source.get("Title_or_Franchise", source_id),
                "x": round(x, 5),
                "y": round(y, 5),
                "size": round(4.2 + min(7.5, math.log2(count + 1) * 1.25), 3),
                "color": COVERAGE_COLORS[level],
                "coverage": level,
                "medium": source.get("Medium", ""),
                "region": source.get("Region_Tradition", ""),
                "priorityTier": source.get("Priority_Tier", ""),
                "entryCount": count,
                "target": target,
                "record": source_record,
            }
        )

    for (source_id, archetype_id), count in sorted(source_archetype_counts.items()):
        graph_edges.append(
            {
                "id": f"mapping:{source_id}:{archetype_id}",
                "source": source_id,
                "target": archetype_id,
                "kind": "mapping",
                "weight": count,
                "label": f"{count} source entr{'y' if count == 1 else 'ies'}",
            }
        )

    for row in entries:
        entry_id = row["Entry_ID"]
        source_id = row["Source_ID"]
        primary_id = row.get("Primary_Archetype_ID", "")
        sx, sy = source_positions[source_id]
        ax, ay = archetype_positions.get(primary_id, (sx, sy))
        angle = 2 * math.pi * hash_fraction(entry_id, "entry")
        x = (sx + ax) / 2 + 1.4 * math.cos(angle)
        y = (sy + ay) / 2 + 1.4 * math.sin(angle)
        entry_record = dict(row)
        entry_record["Terms"] = terms_by_entry.get(entry_id, [])
        graph_nodes.append(
            {
                "id": entry_id,
                "kind": "entry",
                "label": row.get("Canonical_Name", entry_id),
                "x": round(x, 5),
                "y": round(y, 5),
                "size": 2.2,
                "color": "#f8fafc",
                "sourceId": source_id,
                "domain": archetype_by_id.get(primary_id, {}).get("Domain", ""),
                "record": entry_record,
            }
        )
        graph_edges.append(
            {
                "id": f"contains:{source_id}:{entry_id}",
                "source": source_id,
                "target": entry_id,
                "kind": "contains",
                "weight": 1,
                "label": "contains entry",
            }
        )
        for index, archetype_id in enumerate(entry_archetypes[entry_id]):
            graph_edges.append(
                {
                    "id": f"entry-map:{entry_id}:{archetype_id}",
                    "source": entry_id,
                    "target": archetype_id,
                    "kind": "entry-mapping",
                    "weight": 1,
                    "label": "primary mapping" if index == 0 else "additional mapping",
                }
            )

    for row in crosswalks:
        crosswalk_id = row["Crosswalk_ID"]
        adaptation_source_id = row.get("Adaptation_Source_ID", "")
        traditional_source_id = row.get("Traditional_Source_ID", "")
        if not adaptation_source_id or adaptation_source_id not in source_positions:
            continue
        sx, sy = source_positions[adaptation_source_id]
        angle = 2 * math.pi * hash_fraction(crosswalk_id, "crosswalk")
        radius = 5 + 3 * hash_fraction(crosswalk_id, "crosswalk-radius")
        graph_nodes.append(
            {
                "id": crosswalk_id,
                "kind": "crosswalk",
                "label": row.get("Adapted_Concept", crosswalk_id),
                "x": round(sx + radius * math.cos(angle), 5),
                "y": round(sy + radius * math.sin(angle), 5),
                "size": 3.0,
                "color": "#e879f9",
                "operation": row.get("Adaptation_Operation", ""),
                "record": row,
            }
        )
        graph_edges.append(
            {
                "id": f"adapted:{adaptation_source_id}:{crosswalk_id}",
                "source": adaptation_source_id,
                "target": crosswalk_id,
                "kind": "adaptation",
                "weight": 1,
                "label": row.get("Adaptation_Operation", "adapted as"),
            }
        )
        if traditional_source_id in source_by_id:
            graph_edges.append(
                {
                    "id": f"traditional:{traditional_source_id}:{crosswalk_id}",
                    "source": traditional_source_id,
                    "target": crosswalk_id,
                    "kind": "traditional",
                    "weight": 1,
                    "label": "traditional source",
                }
            )

    for row in concepts:
        concept_id = row["Concept_ID"]
        source_id = row["Source_ID"]
        if source_id not in source_positions:
            errors.append(f"{concept_id} has unknown source {source_id}")
            continue
        archetype_ids = split_ids(row.get("Archetype_IDs", ""))
        for archetype_id in archetype_ids:
            if archetype_id not in archetype_by_id:
                errors.append(f"{concept_id} has unknown archetype {archetype_id}")
        sx, sy = source_positions[source_id]
        if archetype_ids:
            ax = sum(archetype_positions[item][0] for item in archetype_ids) / len(archetype_ids)
            ay = sum(archetype_positions[item][1] for item in archetype_ids) / len(archetype_ids)
        else:
            ax, ay = sx, sy
        angle = 2 * math.pi * hash_fraction(concept_id, "concept")
        x = sx * 0.74 + ax * 0.26 + 1.8 * math.cos(angle)
        y = sy * 0.74 + ay * 0.26 + 1.8 * math.sin(angle)
        graph_nodes.append(
            {
                "id": concept_id,
                "kind": "concept",
                "label": row.get("Orientation_Concept", concept_id),
                "x": round(x, 5),
                "y": round(y, 5),
                "size": 2.35,
                "color": "#94a3b8",
                "sourceId": source_id,
                "record": row,
            }
        )
        graph_edges.append(
            {
                "id": f"orientation:{source_id}:{concept_id}",
                "source": source_id,
                "target": concept_id,
                "kind": "orientation",
                "weight": 1,
                "label": "orientation concept",
            }
        )
        for archetype_id in archetype_ids:
            graph_edges.append(
                {
                    "id": f"concept-map:{concept_id}:{archetype_id}",
                    "source": concept_id,
                    "target": archetype_id,
                    "kind": "concept-mapping",
                    "weight": 1,
                    "label": "orientation maps to",
                }
            )

    for row in source_relationships:
        from_id = row.get("From_Source_ID", "")
        to_id = row.get("To_Source_ID", "")
        relationship_id = row.get("Relationship_ID", "")
        if from_id not in source_by_id or to_id not in source_by_id:
            errors.append(f"{relationship_id} has unknown source endpoint {from_id} -> {to_id}")
            continue
        graph_edges.append(
            {
                "id": f"source-rel:{relationship_id}",
                "source": from_id,
                "target": to_id,
                "kind": "source-relationship",
                "weight": 1,
                "label": row.get("Relationship", "related source"),
                "record": row,
            }
        )

    if errors:
        raise SystemExit("Workbook validation failed:\n- " + "\n- ".join(errors[:100]))

    media_counts = Counter(str(source.get("Medium", "Unspecified")) for source in sources)
    coverage_counts = Counter(item["Coverage_Level"] for item in source_coverage)
    domain_counts = Counter(str(row.get("Domain", "Unspecified")) for row in taxonomy)

    payload = {
        "meta": {
            "title": "Fantasy & High-Fantasy Comparative Archetype Atlas",
            "version": "3.0 graph preview",
            "sourceWorkbook": args.input.name,
            "sourceFingerprint": source_fingerprint([args.input, Path(__file__)]),
            "counts": {
                "sources": len(sources),
                "coveredSources": sum(1 for item in source_coverage if item["Entry_Count"] > 0),
                "unchartedSources": sum(1 for item in source_coverage if item["Entry_Count"] == 0),
                "archetypes": len(taxonomy),
                "entries": len(entries),
                "orientationConcepts": len(concepts),
                "orientationCoveredSources": len({row.get("Source_ID") for row in concepts}),
                "sourceRelationships": len(source_relationships),
                "crosswalks": len(crosswalks),
                "nodes": len(graph_nodes),
                "edges": len(graph_edges),
            },
            "coverageCounts": dict(coverage_counts),
            "domainCounts": dict(domain_counts),
            "mediaCounts": dict(media_counts.most_common()),
            "validation": {"errors": 0, "status": "passed"},
        },
        "domains": [
            {"name": domain, "color": color, "count": domain_counts[domain]}
            for domain, color in DOMAIN_COLORS.items()
        ],
        "sources": source_coverage,
        "nodes": graph_nodes,
        "edges": graph_edges,
    }

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8"
    )
    print(
        f"Wrote {args.output.relative_to(ROOT)}: "
        f"{len(graph_nodes):,} nodes, {len(graph_edges):,} edges, "
        f"{len(sources):,} sources ({coverage_counts['uncharted']:,} uncharted)."
    )


if __name__ == "__main__":
    main()
