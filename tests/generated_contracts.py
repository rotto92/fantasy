#!/usr/bin/env python3

from pathlib import Path
import json
import sys
import tempfile

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

import reproducible


workbook = load_workbook(ROOT / "fantasy_high_fantasy_archetype_atlas_v4.xlsx", read_only=True, data_only=True)

guide = workbook["Character Viz Guide"]
guide_rows = {
    str(row[0]): str(row[1])
    for row in guide.iter_rows(min_row=2, values_only=True)
    if row[0] and row[1]
}
assert "normalized being/entity or role/vocation concept" in guide_rows["One visual noun"]
assert "named character" not in guide_rows["One visual noun"]
assert "selected normalized concept" in guide_rows["No global edges"]
shipped_views = guide_rows["Four shipped views"]
assert all(view in shipped_views for view in ["Constellations", "Catalogue", "Relations", "Research"])
assert not any(
    abandoned_control in " ".join(guide_rows.values()).casefold()
    for abandoned_control in ["arrange", "nested group", "x/y comparison"]
)

start = workbook["Start Here"]
assert "normalized-concept atlas" in str(start["A3"].value).lower()
assert "characters and source-native terms are cited evidence" in str(start["A3"].value).lower()
assert "normalized-concept" in str(workbook.properties.subject).lower()
assert "concepts as graph nodes" in str(workbook.properties.description).lower()
assert "only characters are visualization nodes" not in str(workbook.properties.description).lower()

dictionary = workbook["Data Dictionary"]
dictionary_rows = [tuple(str(value or "") for value in row) for row in dictionary.iter_rows(values_only=True)]
assert any(
    row[0] == "Character Catalogue"
    and row[1] == "Character_ID"
    and "never become public graph nodes" in row[5]
    for row in dictionary_rows
)

coverage_sheet = workbook["Source Coverage Audit"]
coverage_headers = [str(cell.value or "") for cell in coverage_sheet[1]]
coverage_rows = [
    dict(zip(coverage_headers, row))
    for row in coverage_sheet.iter_rows(min_row=2, values_only=True)
    if str(row[0] or "").startswith("SRC-")
]
atlas = json.loads((ROOT / "generated" / "atlas.json").read_text(encoding="utf-8"))
atlas_sources = {source["Source_ID"]: source for source in atlas["sources"]}
for row in coverage_rows:
    treatment = row["Recommended_Treatment"]
    assert isinstance(treatment, str) and treatment.strip()
    expected = reproducible.target_for_treatment(treatment)
    assert row["Canonical_Target"] == expected
    assert atlas_sources[row["Source_ID"]]["Coverage_Target"] == expected

assert start["B6"].value == reproducible.source_fingerprint([
    ROOT / "fantasy_high_fantasy_archetype_atlas_v3.xlsx",
    ROOT / "public" / "data" / "characters.json",
    ROOT / "scripts" / "build_version4.py",
])
version3_workbook = load_workbook(
    ROOT / "fantasy_high_fantasy_archetype_atlas_v3.xlsx",
    read_only=True,
    data_only=True,
)
assert version3_workbook["Start Here"]["B6"].value == reproducible.source_fingerprint([
    ROOT / "fantasy_high_fantasy_archetype_atlas_v2.xlsx",
    ROOT / "scripts" / "build_version3.py",
])
assert atlas["meta"]["sourceFingerprint"] == reproducible.source_fingerprint([
    ROOT / "fantasy_high_fantasy_archetype_atlas_v3.xlsx",
    ROOT / "scripts" / "extract_atlas.py",
])

review_index = json.loads((ROOT / "research" / "independent_reviews" / "index.json").read_text(encoding="utf-8"))
ledger_names = {
    str(review["ledger"])
    for review in review_index.get("corpus_wide_reviews", [])
    if isinstance(review, dict) and review.get("ledger")
}
for review in review_index.get("source_review_status", {}).values():
    if isinstance(review, dict):
        ledger_names.update(str(ledger) for ledger in review.get("ledgers", []) if ledger)
zero_character_audit = review_index.get("zero_character_audit", {})
if isinstance(zero_character_audit, dict) and zero_character_audit.get("ledger"):
    ledger_names.add(str(zero_character_audit["ledger"]))
research_root = ROOT / "research"
required_bundle_files = ("sources.json", "characters.json", "relationships.json", "source_terms.json")
discovered_bundles = sorted(
    path
    for path in research_root.glob("batch_*/*")
    if path.is_dir() and any((path / filename).exists() for filename in required_bundle_files)
)
expected_research_fingerprint = reproducible.source_fingerprint([
    ROOT / "fantasy_high_fantasy_archetype_atlas_v3.xlsx",
    research_root / "dimensions.json",
    research_root / "independent_reviews" / "index.json",
    *[research_root / "independent_reviews" / name for name in sorted(ledger_names)],
    ROOT / "scripts" / "validate_character_research.py",
    *discovered_bundles,
])
validation_report = json.loads((research_root / "validation_report.json").read_text(encoding="utf-8"))
characters = json.loads((ROOT / "public" / "data" / "characters.json").read_text(encoding="utf-8"))
constellations = json.loads((ROOT / "public" / "data" / "constellations.json").read_text(encoding="utf-8"))
assert validation_report["source_fingerprint"] == expected_research_fingerprint
assert characters["meta"]["sourceFingerprint"] == expected_research_fingerprint

for node in constellations["nodes"]:
    memberships = node.get("evidenceMemberships")
    assert isinstance(memberships, list), node["id"]
    assert len(memberships) == node["evidenceCount"], node["id"]
    membership_sources = {membership["sourceId"] for membership in memberships}
    assert sorted(membership_sources) == node["sourceIds"], node["id"]
    assert len(membership_sources) == node["sourceCount"], node["id"]

ppl_101 = next(node for node in constellations["nodes"] if node["id"] == "PPL-101")
assert len(ppl_101["examples"]) < len(ppl_101["evidenceMemberships"])
assert sum(
    membership["sourceId"] == "SRC-001"
    for membership in ppl_101["evidenceMemberships"]
) == 13

assert constellations["meta"]["counts"]["sourceExamples"] == sum(
    len(node["evidenceMemberships"])
    for node in constellations["nodes"]
)
assert constellations["meta"]["counts"]["representativeExamples"] == sum(
    len(node["examples"])
    for node in constellations["nodes"]
)
ppl_100 = next(node for node in constellations["nodes"] if node["id"] == "PPL-100")
assert ppl_100["evidenceCount"] == 52

node_memberships = {
    node["id"]: {
        (membership["kind"], membership["id"], membership["sourceId"])
        for membership in node["evidenceMemberships"]
    }
    for node in constellations["nodes"]
}
for edge in (edge for edge in constellations["edges"] if edge["kind"] == "affinity"):
    assert len(edge["evidence"]) == edge["weight"], edge["id"]
    assert all(evidence.get("sourceId") for evidence in edge["evidence"]), edge["id"]
    edge_memberships = {
        (evidence["kind"], evidence["id"], evidence["sourceId"])
        for evidence in edge["evidence"]
    }
    assert edge_memberships <= node_memberships[edge["source"]], edge["id"]
    assert edge_memberships <= node_memberships[edge["target"]], edge["id"]

ppl_100_rol_700 = next(
    edge
    for edge in constellations["edges"]
    if edge["kind"] == "affinity" and {edge["source"], edge["target"]} == {"PPL-100", "ROL-700"}
)
assert sum(
    evidence["sourceId"] == "SRC-015"
    for evidence in ppl_100_rol_700["evidence"]
) == 3

with tempfile.TemporaryDirectory(dir=ROOT / "tests") as directory:
    fixture_root = Path(directory)
    input_path = fixture_root / "input.json"
    dependency_path = fixture_root / "semantic_dependency.py"
    input_path.write_text("{}\n", encoding="utf-8")
    dependency_path.write_text("VALUE = 1\n", encoding="utf-8")
    original_module_path = reproducible.__file__
    try:
        reproducible.__file__ = str(dependency_path)
        first_dependency_fingerprint = reproducible.source_fingerprint([input_path])
        dependency_path.write_text("VALUE = 2\n", encoding="utf-8")
        assert reproducible.source_fingerprint([input_path]) != first_dependency_fingerprint
    finally:
        reproducible.__file__ = original_module_path

    bundle_path = fixture_root / "bundle"
    bundle_path.mkdir()
    (bundle_path / "sources.json").write_text("[]\n", encoding="utf-8")
    unquarantined_fingerprint = reproducible.source_fingerprint([bundle_path])
    marker_path = bundle_path / "QUARANTINED.md"
    marker_path.write_text("Quarantined\n", encoding="utf-8")
    quarantined_fingerprint = reproducible.source_fingerprint([bundle_path])
    assert quarantined_fingerprint != unquarantined_fingerprint
    marker_path.unlink()
    assert reproducible.source_fingerprint([bundle_path]) == unquarantined_fingerprint

print("Generated workbook graph contract passed.")
